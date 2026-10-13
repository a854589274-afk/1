"""Round-1 sample: cover (orig P1), two-document comparison table (orig P12), ALT example page B (orig P31).
Edits the ORIGINAL pptx in place (copy), reusing its cover screenshots; deletes the other slides."""
import sys, os, re, zipfile, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from decklib import *
from charts import bar_with_threshold

SRC, OUT, ART = sys.argv[1], sys.argv[2], sys.argv[3]
prs = Presentation(SRC)
S = {i + 1: s for i, s in enumerate(prs.slides)}

# ======================= P1 cover (orig slide 1) =======================
s = S[1]
clear_shapes(s)
set_bg(s, NAVY)
pic = s.shapes.add_picture(os.path.join(ART, "cover_16x9.png"), 0, 0, prs.slide_width, prs.slide_height)
pic.name = "封面插图(无文字)"
send_to_back(s, pic)
add_text(s, 0.85, 1.55, 6.2, 1.9, [{"runs": [("定量检验程序", {})], "space_after": 0},
                                   {"runs": [("结果可比性验证", {})]}],
         size=46, color=WHITE, bold=True, margin=(0, 0, 0, 0), name="主标题", line_spacing=1.05)
add_text(s, 0.85, 3.65, 6.0, 0.95, [[("WS/T 407-2012 & CNAS-GL047:2021", {})], [("解读与我科应用", {})]],
         size=20, color=RGBColor(0xDD, 0xE6, 0xEF), margin=(0, 0, 0, 0), name="副标题", line_spacing=1.15)
add_text(s, 0.85, 5.25, 5.0, 1.2, [[("主讲人　", {"color": RGBColor(0xA9, 0xB9, 0xCC)}), ("哈小琴", {"bold": True})],
                                   [("第940医院检验科", {})],
                                   [("【待填：讲课日期】", {"color": RGBColor(0xF2, 0xB8, 0x8C), "size": 14})]],
         size=16, color=WHITE, margin=(0, 0, 0, 0), name="主讲信息", line_spacing=1.25)
set_notes(s, "【封面】主标题沿用原稿。副标题：WS/T 407-2012 & CNAS-GL047:2021 解读与我科应用。\n"
             "日期为占位，讲课前填写。插图为代码绘制的扁平矢量图，无任何文字/标志。")

# ======================= P10 comparison table (orig slide 12) =======================
s = S[12]
# keep GL047 cover (id 5) and the WS/T 407 header strip (id 14, inside group id 15)
hoist_from_group(s, 15, {14})
clear_shapes(s, keep_ids={5, 14})
gl_pic, ws_pic = shape_by_id(s, 5), shape_by_id(s, 14)
content_chrome(s, "② 怎么比", "两份文件对照：WS/T 407-2012 与 CNAS-GL047:2021",
               "WS/T 407 §1/§4–§6.9　｜　GL047 前言/§1–§8", 10, badge_w=4.6)

colw = [1.05, 5.54, 5.54]
TX, TY = MX, 1.55
rows = [
    ("性质",
     [[("推荐性", {"bold": True}), ("卫生行业标准，现行", {})], [("2012-12-25发布，2013-08-01实施", {})]],
     [[("CNAS", {}), ("认可指南", {"bold": True}), ("（可参考使用）", {})], [("2021-04-25发布实施", {})]]),
    ("适用",
     [[("同一医疗机构内", {"bold": True}), ("；最多", {}), ("10", {"bold": True}), ("个检测系统；重复检测不超过", {}), ("5", {"bold": True}), ("次", {})]],
     [[("相同或不同地点", {"bold": True}), ("的多个经确认的定量检验程序", {})]]),
    ("前提",
     [[("已按EP9-A2和(或)EP15-A2做性能评价/验证", {})], [("最大CV/最小CV<2倍", {"bold": True})]],
     [[("所用检测系统", {}), ("关键性能指标经验证满足要求", {"bold": True})]]),
    ("样本",
     [[("宜首选", {"bold": True}), ("临床标本；用其他物质须验证互通性", {})], [("至少2个浓度（含正常、异常），取m总×(1±20%)", {})]],
     [[("应使用", {"bold": True}), ("患者样本（抗凝剂相同），", {}), ("不少于5例", {"bold": True})], [("覆盖临床适宜区间，重点关注医学决定水平", {})]]),
    ("方法",
     [[("长期CV → 合并CV → 查表A.1定重复次数 → 比对偏差R", {})]],
     [[("≤4个", {"bold": True}), ("：可选参比系统比对法；", {}), (">4个", {"bold": True}), ("：可选均值法", {})], [("7.2.3 可参考 WS/T 407", {})]]),
    ("判定",
     [[("R ≤ 分析质量要求", {"bold": True}), ("（按6.7.2的7级优先顺序确定）", {})]],
     [[("参比法：5份中≥4份符合，或扩样后90%以上符合", {})], [("均值法：R符合判断标准", {})]]),
    ("不通过",
     [[("剔除偏差较大者后重算；必要时校准", {})], [("难以纠正：与临床沟通，", {}), ("采用不同参考区间", {"bold": True})],
      [("和(或)医学决定水平，并在报告单上标示", {"bold": True})]],
     [[("有计量学可比性：可用修正值/修正因子", {})], [("否则：", {}), ("避免发出不同临床意义的结果", {"bold": True}), ("，或与用户讨论；均记录", {})]]),
    ("时机",
     [[("周期性比对", {})], [("第5章8种情形（质控漂移、换批号、维修后等）", {})]],
     [[("启用前全面验证；定期评审；任一要素变更时评估后重新验证", {})]]),
]
gf, tbl = add_table(s, len(rows) + 1, 3, TX, TY, sum(colw), 0.72 + 0.5 * len(rows), colw, name="对照表(原生)")
tbl.rows[0].height = Inches(0.7)
for r in range(1, len(rows) + 1):
    tbl.rows[r].height = Inches(0.4)
HDR_PIC_W = {1: 1.55, 2: 0.72}
hdr = [("维度", None), ("WS/T 407-2012", "卫生行业标准"), ("CNAS-GL047:2021", "CNAS认可指南")]
for c, (main, sub) in enumerate(hdr):
    cell = tbl.cell(0, c)
    paras = [[(main, {"bold": True, "size": 18})]] if sub is None else [[(main, {"bold": True, "size": 18})], [(sub, {"size": 12, "color": RGBColor(0xC9, 0xD4, 0xE0)})]]
    left = 0.1 if c == 0 else HDR_PIC_W[c] + 0.25
    style_cell(cell, paras, color=WHITE, fill=NAVY, margins=(left, 0.03, 0.08, 0.03), line_pt=21)
    cell_borders(cell, None, WHITE if c < 2 else None, None, None, width=1.5)
for r, (dim, ws, gl) in enumerate(rows, start=1):
    style_cell(tbl.cell(r, 0), [[(dim, {"bold": True})]], size=15, color=NAVY, fill=CARD)
    style_cell(tbl.cell(r, 1), ws, size=14, margins=(0.08, 0.03, 0.06, 0.03))
    style_cell(tbl.cell(r, 2), gl, size=14, margins=(0.08, 0.03, 0.06, 0.03))
    for c in range(3):
        cell_borders(tbl.cell(r, c), None, None, None, RULE, width=0.75)
# cover thumbnails sit inside the header row
gl_pic.crop_top, gl_pic.crop_bottom, gl_pic.crop_left, gl_pic.crop_right = 0.06, 0.42, 0.0, 0.0
ws_pic.crop_top, ws_pic.crop_bottom, ws_pic.crop_left, ws_pic.crop_right = 0.0, 0.0, 0.0, 0.0
thumb_h = 0.6
gl_w = thumb_h * (944 * 1.0) / (1392 * (1 - 0.06 - 0.42))
ws_w = thumb_h * 1400 / 474
x_ws = TX + colw[0] + 0.12
x_gl = TX + colw[0] + colw[1] + 0.12
place(ws_pic, x_ws, TY + 0.07, ws_w, thumb_h)
place(gl_pic, x_gl, TY + 0.07, gl_w, thumb_h)
ws_pic.name, gl_pic.name = "WS/T 407封面(原截图)", "GL047封面(原截图)"
# bring thumbnails above the table
for p in (ws_pic, gl_pic):
    sp = s.shapes._spTree; sp.remove(p._element); sp.append(p._element)
HDR_PIC_W[1], HDR_PIC_W[2] = ws_w, gl_w
for c in (1, 2):
    tbl.cell(0, c).margin_left = Inches(HDR_PIC_W[c] + 0.25)

set_notes(s, """【本页要点】两份文件不是同一适用范围：WS/T 407 限同一医疗机构内、≤10个系统；GL047 覆盖相同或不同地点的多个检验程序。原稿中间的红框与双向箭头已删除，以免误读为两者适用范围相同。

【WS/T 407-2012 原文】
1 范围：本标准规定了对定量检验项目的多个检测系统实施结果可比性验证的要求，包括可比性验证方案的使用条件和适用情况、验证方法和程序、验证结果不符合要求的处理措施。本标准适用于同一医疗机构内的临床实验室。
4.1 本指南规定的比对方案仅适用于最多10个检测系统的结果比对，比对物质的重复检测次数不超过5次。超过10个检测系统的结果比对方法和要求，详见 WS/T 406。
4.2 每个检测系统在使用本指南规定的方案实施比对前，应按照CLSI颁布的EP9-A2和(或)EP15-A2文件的要求，进行全面的性能评价或验证。
6.3.1 宜使用临床标本作为首选比对物质；不得不使用其他物质(如室间质评物或其他参考物质)时，应验证比对物质的互通性。
6.4.2 比较不同检测系统不精密度的大小，确定最大CV与最小CV间的差异是否小于2倍。如小于2倍，可使用本指南规定的比对方案；如大于2倍，则应参照CLSI EP9-A2和(或)EP15-A2确认检测系统间的结果可比性。
6.9 维持结果的可比性需以检测系统各质量保证环节的标准化为前提，必要时通过校准改善结果的可比性，即不同检测系统通过结果的数字转换获得结果的一致性；结果不可比且难以纠正时，应与临床进行沟通，采用不同的参考区间和(或)医学决定水平并在检验报告单上明确标示。

【CNAS-GL047:2021 原文】
前言：本文件……是对CNAS-CL02《医学实验室质量和能力认可准则》中“5.6.4 检验结果的可比性”要求涉及的定量检验程序间检验结果可比性验证所做的具体解释和指导……本文件为指南文件，医学实验室可参考使用，也可使用其它适宜方式对定量检验程序间检验结果可比性进行验证，如应用临床数据进行统计评估。
1 范围：本文件适用于申请认可或已获认可的医学实验室在相同或不同地点，使用多个相同或不同的定量检验程序、检测系统、检验方法等（通称“检验程序”）检测同一分析物（检验项目）时，验证检验结果间的可比性……本文件适用于医学实验室应用的经确认的定量检验程序。
6.4 样本：应使用患者样本（若为抗凝样本，应使用相同的抗凝剂）；样本中被测物浓度、活性等应能覆盖临床适宜区间，重点关注医学决定水平。比对的样本数量应不少于5例。
8 相关措施：……对于具有计量学可比性的检测系统结果，实验室可合理应用修正值或修正因子建立检测系统间结果的可比性，并保留相关记录。……实验室应采取有效措施避免向临床发出具有不同临床意义或解释的结果，或与用户讨论不可比结果对临床活动的影响，以及如何正确应用不同测量系统的检测结果（包括结果偏倚情况），并记录。

【版本说明】GL047 前言引用的是 CL02:2012 的条款号 5.6.4；CL02:2023 对应条款为 7.3.7.4，CNAS-CL02-A001:2023 第7.3.7.4条写明“可参考 CNAS-GL047 以及相关国家/行业标准，如 WS/T 406、WS/T 407”。
状态核查：WS/T 407-2012 在全国标准信息公共服务平台为“现行”；CNAS-GL047:2021 在 CNAS《实验室认可规范文件清单》(2026-9-23) 中仍列为有效文件。""")

# ======================= P23 ALT example B (orig slide 31) =======================
s = S[31]
clear_shapes(s)
content_chrome(s, "② 怎么比", "ALT示例②：比对结果与判定", "WS/T 407 附录B.1 表B.3/B.4", 23, badge_w=3.2)

# B.3 raw results
colw3 = [1.2, 1.25, 1.25, 1.25, 1.25]
gf3, t3 = add_table(s, 6, 5, MX, 1.6, sum(colw3), 2.7, colw3, name="表B.3样本检测结果(原生)")
t3.cell(0, 0).merge(t3.cell(1, 0))
t3.cell(0, 1).merge(t3.cell(0, 2))
t3.cell(0, 3).merge(t3.cell(0, 4))
hdr_fill = NAVY
style_cell(t3.cell(0, 0), [[("检测次数", {"bold": True})]], size=15, color=WHITE, fill=hdr_fill, align=PP_ALIGN.CENTER)
style_cell(t3.cell(0, 1), [[("样本1 / (U/L)", {"bold": True})]], size=15, color=WHITE, fill=hdr_fill, align=PP_ALIGN.CENTER)
style_cell(t3.cell(0, 3), [[("样本2 / (U/L)", {"bold": True})]], size=15, color=WHITE, fill=hdr_fill, align=PP_ALIGN.CENTER)
for c, lab in zip(range(1, 5), ["分析仪A", "分析仪B", "分析仪A", "分析仪B"]):
    style_cell(t3.cell(1, c), [[(lab, {"bold": True})]], size=14, color=NAVY, fill=CARD, align=PP_ALIGN.CENTER)
data3 = [("1", "45.4", "44.1", "126.7", "124.5"), ("2", "46.3", "45.9", "125.5", "122.7"),
         ("3", "46.7", "44.6", "126.2", "123.8"), ("均值", "46.1", "44.9", "126.1", "123.7")]
for r, row in enumerate(data3, start=2):
    last = r == 5
    for c, v in enumerate(row):
        style_cell(t3.cell(r, c), [[(v, {"bold": last or c == 0})]], size=16,
                   color=NAVY if (last or c == 0) else INK, fill=CARD if last else WHITE, align=PP_ALIGN.CENTER)
for r in range(6):
    t3.rows[r].height = Inches(0.42)
    for c in range(5):
        cell_borders(t3.cell(r, c), None, None, None, RULE if r >= 1 else None)

# B.4 analysis
colw4 = [2.0, 2.1, 2.1]
gf4, t4 = add_table(s, 4, 3, MX, 4.6, sum(colw4), 1.9, colw4, name="表B.4数据分析(原生)")
rows4 = [("统计结果", "样本1", "样本2"), ("总均值", "45.5 U/L", "124.9 U/L"),
         ("极差", "1.2 U/L", "2.4 U/L"), ("比对偏差 R", "2.64%", "1.92%")]
for r, row in enumerate(rows4):
    for c, v in enumerate(row):
        if r == 0:
            style_cell(t4.cell(r, c), [[(v, {"bold": True})]], size=15, color=WHITE, fill=NAVY, align=PP_ALIGN.CENTER)
        elif r == 3 and c > 0:
            style_cell(t4.cell(r, c), [[(v, {"bold": True})]], size=20, color=ORANGE, fill=WHITE, align=PP_ALIGN.CENTER)
        else:
            style_cell(t4.cell(r, c), [[(v, {"bold": c == 0})]], size=16, color=NAVY if c == 0 else INK,
                       fill=CARD if c == 0 else WHITE, align=PP_ALIGN.CENTER)
        cell_borders(t4.cell(r, c), None, None, None, RULE if r >= 1 else None)
    t4.rows[r].height = Inches(0.46)

# chart: R vs 6 % criterion
CXo, CYo, CW, CH = 7.25, 2.15, 5.48, 3.75
gfc, pb = bar_with_threshold(s, CXo, CYo, CW, CH, ["样本1", "样本2"], (2.64, 1.92),
                             threshold=6, vmax=7, plot_box=(0.11, 0.05, 0.87, 0.8))
badge = add_box(s, CXo + 0.1, 1.6, 1.55, 0.44, fill=TEAL, radius=0.22, name="通过徽章")
box_text(badge, [[("✓ 通过", {"bold": True})]], size=18, color=WHITE, align=PP_ALIGN.CENTER, margin=(0.05, 0, 0.05, 0))
add_text(s, CXo, 6.0, CW, 0.75,
         [[("结论：两个浓度的R均≤6%，A、B两台生化分析仪ALT结果可比", {})]],
         size=18, color=INK, margin=(0.05, 0, 0, 0), name="结论")

set_notes(s, """【WS/T 407 附录B.1.5–B.1.6 原文】样本检测结果见表B.3；数据分析结果见表B.4。B.1.6 结论：两个浓度样本的比对偏差均小于分析质量要求，故认为A和B生化分析仪检测ALT的结果可比性是可接受的。
6.8.5 a) 如比对偏差小于或等于分析质量要求，得出的结论是：在该次评估的样本浓度水平，所有检测系统的结果都具有可比性。

【计算（已用代码复算，正确）】R=(最大均值−最小均值)/总均值×100%。样本1：(46.1−44.9)/45.5=2.64%；样本2：(126.1−123.7)/124.9=1.92%。表中均值为四舍五入到0.1 U/L后的值；若用未取整均值（46.133/44.867；126.133/123.667），R=2.78%/1.98%，同样≤6%，结论不变。

【关于表A.1/表B.2“2.67”】表A.1及表B.2中“2个检测系统、检测3次、CV合并1%”一格印为2.67，与同行其余各列不一致（同行2%列4.53、3%列6.80，折合每1%均为2.27）；按 q0.95(k, k(n−1))/√n × CV合并 复算，该格应为2.27。本例CV合并=2.39%，查的是2%（4.53）与3%（6.80）两列，不受该格影响。页面与引用均保留标准原值2.67，不作修改。

【判定标准来源】6% = 示例中ALT个体内生物学变异CVI 18.0%的1/3（WS/T 407 附录B.1.3，为标准示例取值）。参考：现行 WS/T 403—2024《临床化学检验常用项目分析质量标准》表A.1，ALT 允许不精密度6.0%、允许偏倚5.0%、允许总误差6 U/L(≤40 U/L)/15.0%(>40 U/L)。本科实际采用何种判定标准，由科室按6.7.2顺序确定（6.7.2 g：所选定的分析质量要求至少应满足国家或行业标准的要求）。

【图表说明】原生图表：柱=表B.4的R值，虚线=判定标准6%（均可在“编辑数据”中修改）。两侧各留一个空白类别，仅用于让判定虚线延伸到柱外。""")

# ======================= drop all other slides =======================
keep = {1, 12, 31}
sldIdLst = prs.slides._sldIdLst
for i, sldId in enumerate(list(sldIdLst), start=1):
    if i not in keep:
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)
prs.save(OUT)

# ======================= theme fonts: Arial + 微软雅黑 =======================
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if re.match(r"ppt/theme/theme\d+\.xml", item.filename):
            x = data.decode("utf-8")
            for blk in ("majorFont", "minorFont"):
                def fix(m):
                    b = m.group(0)
                    b = re.sub(r'<a:latin typeface="[^"]*"', '<a:latin typeface="Arial"', b)
                    b = re.sub(r'<a:ea typeface="[^"]*"', '<a:ea typeface="微软雅黑"', b)
                    b = re.sub(r'<a:font script="Hans" typeface="[^"]*"', '<a:font script="Hans" typeface="微软雅黑"', b)
                    return b
                x = re.sub(rf"<a:{blk}>.*?</a:{blk}>", fix, x, flags=re.S)
            data = x.encode("utf-8")
        zout.writestr(item, data)
shutil.move(tmp, OUT)
print("saved", OUT)
