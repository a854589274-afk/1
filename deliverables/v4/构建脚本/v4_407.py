# -*- coding: utf-8 -*-
"""第4轮：在用户修改后的PPT上操作（不改动用户已有页面的内容与版式）。
1) 新增 WS/T 407 分隔页（仿第18页 GL047 分隔页）；
2) 新增 3 页 WS/T 407 解读：适用范围与使用条件（1、4、5章）；比对项目/仪器/样本（6.1–6.3、3.2）；示例2 RBC 与不通过处理（附录B.2、6.8.6、6.9）；
3) 把原第33–38页（WS/T 407 比对实施与 ALT 示例）连同新页移到 GL047 分隔页（原第18页）之前；
4) 更新全部页脚页码。
用法：python3 v4_407.py 用户PPT.pptx 输出.pptx"""
import sys, os, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from lib import *
from kit import *
from core import img
from pages_a import native_table, WHITE

X0, FW = ML, SW - ML - MR


class _D:
    mode = 'clean'


def new_slide(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    for ph in list(sl.placeholders):
        ph._element.getparent().remove(ph._element)
    return sl


def T(s):
    return P(s)


def divider(prs, ctx, page):
    sl = new_slide(prs)
    picture(sl.shapes, img('407_hero'), 0, 0, SW, SH, rounded=False, name='NEWIMG_HERO')
    textbox(sl.shapes, ML, Y_CHAP, 4.0, H_CHAP, [P(CHAPTERS[2])], 12, color='9FE3DE', bold=True, pitch=1.2, name='CHAPTER')
    # 与第18页 GL047 分隔页相同位置与字号：标题 32pt 三行，副标题 20pt
    textbox(sl.shapes, ML, 2.3, span(7), 1.73, [P('WS/T 407-2012'), P('医疗机构内定量检验结果的'), P('可比性验证指南')], 32,
            color=WHITE, bold=True, pitch=1.2, force=True)
    textbox(sl.shapes, ML, 4.33, span(7), 0.35, [P('2012-12-25发布　2013-08-01实施')], 20, color='9FE3DE', pitch=1.2)
    footer(sl, page, dark=True)
    ctx.page = page
    ctx.L(template='T1 主视觉页（分隔页，仿第18页）', orig='新增页', size='32pt / 20pt（同第18页）', images='407_hero',
          added=['标题「WS/T 407-2012 医疗机构内定量检验结果的可比性验证指南」三行', '副标题「2012-12-25发布　2013-08-01实施」', '主视觉 407_hero'])
    return sl


def page_scope(prs, ctx, page):
    sl = new_slide(prs); ctx.page = page
    cite = lambda c: '—— WS/T 407-2012 ' + c
    items = ['室内质控结果有漂移趋势时。', '室间质评结果不合格，采取纠正措施后。', '更换试剂批号（必要时）。', '更换重要部件或重大维修后。',
             '软件程序变更后。', '临床医生对结果的可比性有疑问时。', '患者投诉对结果可比性有疑问（需要确认时）。', '需提高周期性比对频率时（如每季度或每月一次）。']
    def draw(sz, v):
        frame(sl, [P('WS/T 407 适用范围与使用条件')], 2, page, '来源：WS/T 407-2012 第1、4、5章', 'WS', corner=img('407_scope'), mode='clean')
        ctx.L(images='407_scope')
        cw = {'a': 6.4, 'b': 6.7}[v]
        y = CT
        y += qb(ctx, sl, X0, y, cw, [P('1 范围', bold=True, color=C['TEAL']),
                                       P('本标准规定了对定量检验项目的多个检测系统实施结果可比性验证的要求，包括可比性验证方案的使用条件和适用情况、验证方法和程序、验证结果不符合要求的处理措施。'),
                                       P('本标准适用于', P('同一医疗机构内', bold=True, color=C['ORANGE'])[0], '的临床实验室。')], sz, 'WS', cite('1')) + GUT
        qb(ctx, sl, X0, y, cw, [P('4 可比性验证方案的使用条件', bold=True, color=C['TEAL']),
                                 P('4.1 本指南规定的比对方案仅适用于', P('最多10个检测系统', bold=True, color=C['ORANGE'])[0], '的结果比对，比对物质的重复检测次数', P('不超过5次', bold=True, color=C['ORANGE'])[0], '。超过10个检测系统的结果比对方法和要求，详见 WS/T 406。'),
                                 P('4.2 每个检测系统在使用本指南规定的方案实施比对前，应按照美国临床和实验室标准协会（Clinical and Laboratory Standards Institute，CLSI）颁布的 EP9-A2 和（或）EP15-A2 文件的要求，进行全面的性能评价或验证。'),
                                 P('4.3 可比性验证只能确认检测系统性能的一个方面，不能取代其他质量保证环节，如校准和室内质量控制等。')], sz, 'WS', cite('4'))
        # 右栏：第5章 8 种适用情况（编号卡 2 列 × 4 行）
        rx = X0 + cw + GUT; cw = FW - cw - GUT
        h = text(ctx, sl, rx, CT, cw, [P('5 使用本指南规定的方案进行可比性验证的适用情况', bold=True, color=C['TEAL'])], sz, limit=CB_R)
        yy = CT + h + 0.1
        ccw = (cw - 0.12) / 2
        hs = [text_height_in([P(t)], ccw - 0.62, sz) + 0.16 for t in items]
        for r in range(4):
            rh = max(hs[2 * r], hs[2 * r + 1])
            for c in range(2):
                k = 2 * r + c
                x = rx + c * (ccw + 0.12)
                g = group(sl.shapes, 'NUMCARD')
                rrect(g.shapes, x, yy, ccw, rh, fill=WHITE, line=C['TEAL'])
                textbox(g.shapes, x + 0.08, yy + 0.08, 0.5, 0.3, [P(f'5.{k+1}')], 14, color=C['ORANGE'], bold=True, pitch=1.15)
                textbox(g.shapes, x + 0.56, yy + 0.08, ccw - 0.64, rh - 0.16, [P(items[k])], sz, pitch=PITCH)
            yy += rh + 0.1
        yy += 0.05
        caption(sl.shapes, rx, yy, cw, '—— WS/T 407-2012 5', align='r', color=C['TEAL'])
        yy += 0.3
        ctx.chk(yy, CB_R)
    sz = fit(ctx, sl, draw, [18, 17, 16, 15, 14], ('a', 'b'))
    ctx.L(template='T2 条文页（通栏两列，新图在角图位）', orig='新增页（WS/T 407 解读①）', size=f'正文{sz}pt',
          added=['第1章范围原文（突出"同一医疗机构内"）', '第4章4.1–4.3原文（突出"最多10个检测系统""不超过5次"）', '第5章8种适用情况编号卡', '新图 407_scope（角图）'],
          pending=['4.1所引WS/T 406为2012版，现行为WS/T 406—2024（标题已改；上轮书目已核）——放不下，只写入日志，讲时可口头补充'])
    return sl


def page_sample(prs, ctx, page):
    sl = new_slide(prs); ctx.page = page
    cite = lambda c: '—— WS/T 407-2012 ' + c
    def draw(sz, v):
        frame(sl, [P('比对项目、比对仪器与比对样本')], 2, page, '来源：WS/T 407-2012 3.2、6.1–6.3', 'WS', corner=img('407_sample'), mode='clean')
        ctx.L(images='407_sample')
        bl = [QB([P('6.1 确定需要进行结果比对的检测项目', bold=True, color=C['TEAL']),
                  P('同一医疗机构内，使用多个检测系统向临床报告检验结果的检测项目', P('均应进行可比性验证', bold=True, color=C['ORANGE'])[0], '。')], 'WS', cite('6.1')),
              QB([P('6.2 确定需要进行结果比对的检测仪器', bold=True, color=C['TEAL']),
                  P('实施周期性结果比对时，用于检测临床样本的', P('每个检测系统', bold=True, color=C['ORANGE'])[0], '都应进行可比性验证。')], 'WS', cite('6.2')),
              QB([P('6.3 比对样本　6.3.1 样本选择', bold=True, color=C['TEAL']),
                  P(P('宜使用临床标本作为首选比对物质', bold=True, color=C['ORANGE'])[0], '；不得不使用其他物质（如室间质评物或其他参考物质）时，应验证比对物质的互通性。')], 'WS', cite('6.3.1')),
              QB([P('6.3.2 样本浓度水平', bold=True, color=C['TEAL']),
                  P('使用本指南规定的比对方案，需已知比对物质不同浓度水平对应的结果不精密度，故通常选择与质控品浓度水平相近的比对物质进行可比性验证；要求每个检测系统', P('至少检测两个浓度水平（包含正常和异常水平）', bold=True, color=C['ORANGE'])[0], '的比对物质。')], 'WS', cite('6.3.2')),
              QB([P('3.2 互通性 commutability（互换性；替换性）', bold=True, color=C['TEAL']),
                  P('用不同测量程序测量该物质时，各测量程序所得测量结果之间的数字关系，与用这些测量程序测量实际临床标本时测量结果的数字关系的一致程度。')], 'WS', cite('3.2'))]
        def bh(w, sz): return 1.25
        def bd(ctx, sl, x, y, w, sz, limit):
            hw = (w - GUT) / 2
            h1 = bignum(sl.shapes, x, y, hw, '≥ 2', '浓度水平（正常＋异常）', nsz=34)
            h2 = bignum(sl.shapes, x + hw + GUT, y, hw, '±20%', '以 m总 为中心选样本浓度（6.5）', nsz=34)
            ctx.chk(y + max(h1, h2), limit)
            return max(h1, h2)
        pack(ctx, sl, bl, X0, FW, 2, sz)
    sz = fit(ctx, sl, draw, [18, 17, 16, 15, 14], ('a',))
    ctx.L(template='T2 条文页（通栏两列，新图在角图位）', orig='新增页（WS/T 407 解读②）', size=f'正文{sz}pt',
          added=['6.1、6.2、6.3.1、6.3.2原文条文框（关键词橙色）', '3.2互通性定义', '新图 407_sample（角图）'])
    return sl


def page_rbc(prs, ctx, page):
    sl = new_slide(prs); ctx.page = page
    cite = lambda c: '—— WS/T 407-2012 ' + c
    def draw(sz, v):
        frame(sl, [P('示例2：4台血液分析仪RBC——不通过怎么办')], 2, page, '来源：WS/T 407-2012 附录B.2、6.8.6、6.9', 'WS', corner=img('407_rbc'), mode='clean')
        ctx.L(images='407_rbc')
        lw = {'a': 6.2, 'b': 6.0, 'c': 5.8}[v]
        y = CT
        y += text(ctx, sl, X0, y, lw, [P('质量要求（B.2.3）：', P('低浓度6%（RBC TEa）；正常、高浓度3%（1/2 TEa）', bold=True, color=C['ORANGE'])[0], '。以CVI（3.2%）的1/3（1.1%）为要求"过于严格"，未采用。')], sz) + 0.15
        heads = [[P(t)] for t in ['样本', '要求', '重复', '总均值', '极差', 'R', '结论']]
        R = C['RED']; G = C['GREEN']
        rows = [
            [[P('样本1', bold=True)], [P('6%')], [P('5次')], [P('2.32')], [P('0.19')], [P('8.19%', bold=True, color=R)], [P('未通过', bold=True, color=R)]],
            [[P('样本1 剔除B', bold=True)], [P('6%')], [P('—')], [P('2.28')], [P('0.07')], [P('3.07%', bold=True, color=G)], [P('通过', bold=True, color=G)]],
            [[P('样本2', bold=True)], [P('3%')], [P('3次')], [P('4.46')], [P('0.12')], [P('2.69%', bold=True, color=G)], [P('通过', bold=True, color=G)]],
            [[P('样本3', bold=True)], [P('3%')], [P('3次')], [P('5.94')], [P('0.10')], [P('1.68%', bold=True, color=G)], [P('通过', bold=True, color=G)]],
        ]
        widths = [1.25, 0.65, 0.65, 0.85, 0.75, 0.85, 0.95]
        widths[0] += lw - sum(widths)
        tot, _, _ = native_table(ctx, sl, X0, y, widths, heads, rows, 14, [C['TEAL']] * 7, limit=CB_L)
        y += tot + 0.08
        y += text(ctx, sl, X0, y, lw, [P('单位：×10¹²/L；数据取自表B.7–B.9（本轮已复算）', color=C['GRAY'])], 12) + 0.12
        text(ctx, sl, X0, y, lw, [P('样本1：差异最大的B与D各与规范操作系统A比：', P('B与A偏差6.11%，D与A偏差2.18%', bold=True, color=C['ORANGE'])[0], ' → 剔除B，A、C、D重算R＝3.07%，通过（B.2.5）。')], sz)
        rx = X0 + lw + GUT; rw = FW - lw - GUT
        yy = CT
        yy += qb(ctx, sl, rx, yy, rw, [P(P('6.8.6 ', bold=True, color=C['TEAL'])[0], '对于不符合可比性要求的检测系统，应分析原因，必要时采取相应的纠正措施，其后再将该检测系统与规范操作检测系统（如使用配套试剂、用配套校准物定期进行仪器校准、仪器性能良好、规范地开展室内质量控制、参加室间质量评价成绩优良、检测程序规范、人员经过良好培训的检测系统）的结果进行比对，确认比对结果符合分析质量要求。')],
                 sz, 'WS', cite('6.8.6'), limit=CB_R) + GUT
        qb(ctx, sl, rx, yy, rw, [P('6.9 可比性验证结果不符合要求的处理措施', bold=True, color=C['TEAL']),
                                 P('维持结果的可比性需以检测系统各质量保证环节的标准化为前提，必要时通过校准改善结果的可比性，即不同检测系统通过结果的数字转换获得结果的一致性；结果不可比且难以纠正时，应与临床进行沟通，', P('采用不同的参考区间和（或）医学决定水平并在检验报告单上明确标示', bold=True, color=C['ORANGE'])[0], '。')],
           sz, 'WS', cite('6.9'), limit=CB_R)
    sz = fit(ctx, sl, draw, [18, 17, 16, 15, 14], ('a', 'b', 'c'))
    ctx.L(template='T5 数据页（左：标准与原生结果表；右：不通过处理条文）', orig='新增页（WS/T 407 解读③）', size=f'正文{sz}pt',
          added=['B.2.3可接受标准摘要（一行，非原文直引；原文放不下）', '原生结果表：样本1 R=8.19%未通过→剔除B后3.07%通过；样本2 2.69%、样本3 1.68%通过', 'B.2.5剔除过程摘要（一句，非原文直引）', '6.8.6、6.9原文条文框', '新图 407_rbc（角图）'],
          pending=['水平2质控 CV最大/最小 = 2.01/0.86 = 2.34，已超过6.4.2"小于2倍"的前提，标准示例仍按本方案进行（复算发现，只写入日志，未上页面）'])
    return sl


def renumber(prs):
    for i, sl in enumerate(prs.slides, 1):
        foots = [sh for sh in sl.shapes if sh.name == 'FOOTER' and sh.has_text_frame and sh.text_frame.text.strip().isdigit()]
        for sh in foots:
            runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
            runs[0].text = str(i)
            for r in runs[1:]: r.text = ''


if __name__ == '__main__':
    src, dst = sys.argv[1], sys.argv[2]
    prs = Presentation(src)
    n0 = len(prs.slides)
    assert n0 == 40
    old = list(prs.slides._sldIdLst)
    ctx = Ctx(_D())
    # 终稿页码：1–17 不变；18 分隔；19、20 新页；21–26 原33–38；27 新页；28–42 原18–32；43–44 原39–40
    s_div = divider(prs, ctx, 18)
    s_a = page_scope(prs, ctx, 19)
    s_b = page_sample(prs, ctx, 20)
    s_c = page_rbc(prs, ctx, 27)
    new = list(prs.slides._sldIdLst)[n0:]
    lst = prs.slides._sldIdLst
    order = old[:17] + [new[0], new[1], new[2]] + old[32:38] + [new[3]] + old[17:32] + old[38:40]
    for el in list(lst): lst.remove(el)
    for el in order: lst.append(el)
    renumber(prs)
    prs.save(dst)
    import json
    json.dump(dict(notes=ctx.notes), open(os.path.splitext(dst)[0] + '_v4log.json', 'w'), ensure_ascii=False, indent=1, default=str)
    print('slides', len(prs.slides))
