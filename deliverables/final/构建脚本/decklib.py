"""Shared helpers for rebuilding the comparability deck on top of the original pptx."""
import copy
from lxml import etree
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x14, 0x8F, 0x8A)
ORANGE = RGBColor(0xE8, 0x83, 0x3A)
RED = RGBColor(0xC0, 0x39, 0x2B)
CARD = RGBColor(0xF3, 0xF5, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x26, 0x32, 0x40)      # body text
MUTED = RGBColor(0x6B, 0x77, 0x85)    # footer / captions
RULE = RGBColor(0xD9, 0xDE, 0xE4)     # table hairlines

EA_FONT = "微软雅黑"
LATIN_FONT = "Arial"

SLIDE_W = 13.333
MX = 0.6                 # left/right margin (in)
TAG_Y, TAG_H = 0.28, 0.32
TITLE_Y, TITLE_H = 0.66, 0.72
CONTENT_Y = 1.6
FOOT_Y = 7.0

_RPR_AFTER_FONTS = {qn("a:sym"), qn("a:hlinkClick"), qn("a:hlinkMouseOver"), qn("a:rtl"), qn("a:extLst")}


def _set_fonts(rPr, ea=EA_FONT, latin=LATIN_FONT):
    for tag in ("a:latin", "a:ea", "a:cs"):
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)
    new = []
    for tag, face in (("a:latin", latin), ("a:ea", ea), ("a:cs", latin)):
        el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", face)
        rPr.remove(el)
        new.append(el)
    # insert before the first element that must follow the font elements
    anchor = None
    for child in rPr:
        if child.tag in _RPR_AFTER_FONTS:
            anchor = child
            break
    for el in new:
        if anchor is not None:
            anchor.addprevious(el)
        else:
            rPr.append(el)


def style_run(run, size=None, bold=None, color=None, italic=None):
    rPr = run._r.get_or_add_rPr()
    if size is not None:
        rPr.set("sz", str(int(round(size * 100))))
    if bold is not None:
        rPr.set("b", "1" if bold else "0")
    if italic is not None:
        rPr.set("i", "1" if italic else "0")
    rPr.set("lang", "zh-CN")
    rPr.set("altLang", "en-US")
    if color is not None:
        run.font.color.rgb = color
    _set_fonts(rPr)


def fill_text(tf, paras, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
              line_spacing=None, space_after=None):
    """paras: list of paragraphs; each paragraph is a str or a list of (text, opts) runs.
    Paragraph-level overrides can be given as dict {'runs': [...], 'align':..., 'space_after':..., 'size':...}."""
    tf.clear()
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        popts = {}
        if isinstance(para, dict):
            popts = para
            runs = para["runs"]
        elif isinstance(para, str):
            runs = [(para, {})]
        else:
            runs = para
        p.alignment = popts.get("align", align)
        ls = popts.get("line_spacing", line_spacing)
        if ls is not None:
            p.line_spacing = ls
        sa = popts.get("space_after", space_after)
        if sa is not None:
            p.space_after = Pt(sa)
        for text, ro in runs:
            r = p.add_run()
            r.text = text
            style_run(r, size=ro.get("size", popts.get("size", size)),
                      bold=ro.get("bold", bold), color=ro.get("color", color),
                      italic=ro.get("italic"))
    return tf


def add_text(slide, x, y, w, h, paras, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, margin=(0.05, 0.03, 0.05, 0.03), name=None,
             line_spacing=None, space_after=None, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margin]
    fill_text(tf, paras, size=size, color=color, bold=bold, align=align,
              line_spacing=line_spacing, space_after=space_after)
    return tb


def add_box(slide, x, y, w, h, fill=CARD, line=None, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            name=None, line_w=1.0, dash=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = dash
    shp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        # adj is fraction of the shorter side
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    return shp


def box_text(shp, paras, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
             margin=None):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    if margin:
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margin]
    fill_text(tf, paras, size=size, color=color, bold=bold, align=align)
    return shp


def clear_shapes(slide, keep_ids=()):
    sp_tree = slide.shapes._spTree
    for shp in list(slide.shapes):
        if shp.shape_id in keep_ids:
            continue
        sp_tree.remove(shp._element)


def hoist_from_group(slide, group_id, pic_ids):
    """Move pictures out of a group into the slide tree (same slide part, so r:embed stays valid).
    Group-child coordinates are converted with the group's transform."""
    sp_tree = slide.shapes._spTree
    grp = next(s for s in slide.shapes if s.shape_id == group_id)
    gx = grp._element.grpSpPr.find(qn("a:xfrm"))
    off, ext = gx.find(qn("a:off")), gx.find(qn("a:ext"))
    choff, chext = gx.find(qn("a:chOff")), gx.find(qn("a:chExt"))
    sx = int(ext.get("cx")) / int(chext.get("cx"))
    sy = int(ext.get("cy")) / int(chext.get("cy"))
    moved = []
    for child in grp.shapes:
        if child.shape_id in pic_ids:
            el = copy.deepcopy(child._element)
            x = el.find(qn("p:spPr")).find(qn("a:xfrm"))
            o, e = x.find(qn("a:off")), x.find(qn("a:ext"))
            nx = int(off.get("x")) + (int(o.get("x")) - int(choff.get("x"))) * sx
            ny = int(off.get("y")) + (int(o.get("y")) - int(choff.get("y"))) * sy
            o.set("x", str(int(nx))); o.set("y", str(int(ny)))
            e.set("cx", str(int(int(e.get("cx")) * sx))); e.set("cy", str(int(int(e.get("cy")) * sy)))
            sp_tree.append(el)
            moved.append(child.shape_id)
    sp_tree.remove(grp._element)
    return moved


def shape_by_id(slide, sid):
    return next(s for s in slide.shapes if s.shape_id == sid)


def place(shape, x, y, w=None, h=None):
    shape.left, shape.top = Inches(x), Inches(y)
    if w is not None:
        shape.width = Inches(w)
    if h is not None:
        shape.height = Inches(h)


def send_to_back(slide, shape):
    sp_tree = slide.shapes._spTree
    sp_tree.remove(shape._element)
    sp_tree.insert(2, shape._element)


def set_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


# ---------- page chrome ----------

def chapter_tag(slide, text):
    shp = add_box(slide, MX, TAG_Y, 1.55, TAG_H, fill=CARD, radius=0.16, name="章节标签")
    box_text(shp, [text], size=12, color=TEAL, bold=True, align=PP_ALIGN.LEFT,
             margin=(0.12, 0.02, 0.08, 0.02))
    return shp


def title(slide, text, color=NAVY, size=30):
    return add_text(slide, MX, TITLE_Y, SLIDE_W - 2 * MX, TITLE_H, [text], size=size, color=color,
                    bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0, 0, 0), name="标题", wrap=False)


def source_badge(slide, text, width=None):
    w = width or max(1.6, 0.13 * len(text) + 0.35)
    x = SLIDE_W - MX - w
    shp = add_box(slide, x, FOOT_Y - 0.03, w, 0.34, fill=CARD, radius=0.12, name="来源徽章")
    box_text(shp, [text], size=12, color=NAVY, align=PP_ALIGN.CENTER, margin=(0.08, 0.01, 0.08, 0.01))
    return shp


def footer(slide, page_no, dark=False):
    color = RGBColor(0xC9, 0xD4, 0xE0) if dark else MUTED
    return add_text(slide, MX, FOOT_Y, 4.5, 0.3, [[(str(page_no), {"bold": True}), ("　｜　第940医院检验科", {})]],
                    size=12, color=color, anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0, 0, 0), name="页脚")


def content_chrome(slide, chapter, title_text, source, page_no, badge_w=None):
    set_bg(slide, WHITE)
    chapter_tag(slide, chapter)
    title(slide, title_text)
    source_badge(slide, source, badge_w)
    footer(slide, page_no)


def set_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = text


# ---------- tables ----------

def _ln(tag, color, width_pt):
    ln = etree.Element(qn(tag))
    ln.set("w", str(int(width_pt * 12700)))
    ln.set("cap", "flat"); ln.set("cmpd", "sng"); ln.set("algn", "ctr")
    if color is None:
        etree.SubElement(ln, qn("a:noFill"))
    else:
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", str(color))
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
    return ln


def cell_borders(cell, left=None, right=None, top=None, bottom=None, width=0.75):
    """Each side: RGBColor or None (= no line)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    els = [_ln("a:lnL", left, width), _ln("a:lnR", right, width), _ln("a:lnT", top, width),
           _ln("a:lnB", bottom, width)]
    for el in reversed(els):
        tcPr.insert(0, el)


def style_cell(cell, paras, size=14, color=INK, bold=False, fill=WHITE, align=PP_ALIGN.LEFT,
               anchor=MSO_ANCHOR.MIDDLE, margins=(0.08, 0.04, 0.08, 0.04), line_pt=None):
    """line_pt: exact line spacing in points (font-independent row heights); default = 1.2 x size."""
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = anchor
    cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = [Inches(m) for m in margins]
    tf = cell.text_frame
    tf.word_wrap = True
    fill_text(tf, paras, size=size, color=color, bold=bold, align=align,
              line_spacing=Pt(line_pt if line_pt else round(size * 1.22)))


NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def add_table(slide, rows, cols, x, y, w, h, col_widths, name=None):
    gf = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        gf.name = name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE_TABLE
    tbl.first_row = False
    tbl.horz_banding = False
    for i, cw in enumerate(col_widths):
        tbl.columns[i].width = Inches(cw)
    return gf, tbl
