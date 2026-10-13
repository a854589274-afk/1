"""Extra helpers for the full deck: slide-number footer, arrows, icons, equations, doughnut."""
import os, io, uuid
from lxml import etree
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from decklib import *
from decklib import _set_fonts

ICON_DIR = None   # set by builder


def footer_field(slide, dark=False):
    color = RGBColor(0xC9, 0xD4, 0xE0) if dark else MUTED
    tb = add_text(slide, MX, FOOT_Y, 4.5, 0.3, [[("#", {"bold": True}), ("　｜　第940医院检验科", {})]],
                  size=12, color=color, anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0, 0, 0), name="页脚")
    p = tb.text_frame.paragraphs[0]
    r0 = p.runs[0]._r
    fld = etree.Element(qn("a:fld"))
    fld.set("id", "{" + str(uuid.uuid4()).upper() + "}")
    fld.set("type", "slidenum")
    rPr = r0.find(qn("a:rPr"))
    fld.append(rPr)
    t = etree.SubElement(fld, qn("a:t")); t.text = "‹#›"
    r0.addprevious(fld)
    p._p.remove(r0)
    return tb


def chrome(slide, chapter, title_text, source, badge_w=None, title_size=30):
    set_bg(slide, WHITE)
    chapter_tag(slide, chapter)
    title(slide, title_text, size=title_size)
    source_badge(slide, source, badge_w)
    footer_field(slide)


def arrow(slide, x1, y1, x2, y2, color=NAVY, width=2.0, dash=None, head=True, name="箭头"):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.name = name
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if head:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def icon(slide, name, x, y, size):
    pic = slide.shapes.add_picture(os.path.join(ICON_DIR, name + ".png"), Inches(x), Inches(y), Inches(size), Inches(size))
    pic.name = "图标-" + name
    return pic


def picture_from_blob(slide, blob, x, y, w=None, h=None, name=None):
    pic = slide.shapes.add_picture(io.BytesIO(blob), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    if name:
        pic.name = name
    return pic


def num_circle(slide, x, y, d, text, fill=TEAL, color=WHITE, size=18):
    shp = add_box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL, radius=None, name="序号")
    box_text(shp, [[(text, {"bold": True})]], size=size, color=color, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    return shp


def card(slide, x, y, w, h, fill=CARD, line=None, radius=0.12, name="卡片", line_w=1.25):
    return add_box(slide, x, y, w, h, fill=fill, line=line, radius=radius, name=name, line_w=line_w)


def card_text(slide, x, y, w, h, paras, fill=CARD, line=None, size=16, color=INK, anchor=MSO_ANCHOR.TOP,
              margin=(0.16, 0.12, 0.14, 0.1), align=PP_ALIGN.LEFT, name="卡片", radius=0.12, line_spacing=None):
    shp = card(slide, x, y, w, h, fill=fill, line=line, name=name, radius=radius)
    box_text(shp, paras, size=size, color=color, align=align, anchor=anchor, margin=margin)
    if line_spacing:
        for p in shp.text_frame.paragraphs:
            p.line_spacing = line_spacing
    return shp


# ---------------- OMML equation with plain-text fallback ----------------
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def m_r(text, sz, italic=True, color="1F3A5F"):
    sty = "" if italic else '<m:rPr><m:sty m:val="p"/></m:rPr>'
    i = "1" if italic else "0"
    t = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
    return (f'<m:r>{sty}<a:rPr lang="en-US" sz="{sz}" i="{i}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="Cambria Math"/><a:ea typeface="微软雅黑"/></a:rPr><m:t>{t}</m:t></m:r>')


def m_sub(base, sub, sz):
    return f'<m:sSub><m:e>{m_r(base, sz)}</m:e><m:sub>{m_r(sub, sz, italic=False)}</m:sub></m:sSub>'


def m_bar(inner):
    return f'<m:bar><m:barPr><m:pos m:val="top"/></m:barPr><m:e>{inner}</m:e></m:bar>'


def m_frac(num, den):
    return f'<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>'


def add_equation(slide, x, y, w, h, omml, fallback, size=28, name="公式"):
    sid = max([s.shape_id for s in slide.shapes] + [1]) + 1
    sz = int(size * 100)
    X, Y, W, H = [int(Inches(v)) for v in (x, y, w, h)]
    sppr = (f'<p:spPr><a:xfrm><a:off x="{X}" y="{Y}"/><a:ext cx="{W}" cy="{H}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>')
    nv = f'<p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
    body = '<a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"><a:noAutofit/></a:bodyPr><a:lstStyle/>'
    fb = fallback.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    xml = (f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
           f'xmlns:p="{P_NS}" xmlns:a="{A_NS}">'
           f'<mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14">'
           f'<p:sp>{nv}{sppr}<p:txBody>{body}<a:p><a:pPr algn="l"/><a14:m>'
           f'<m:oMathPara xmlns:m="{M_NS}"><m:oMathParaPr><m:jc m:val="left"/></m:oMathParaPr><m:oMath>{omml}</m:oMath></m:oMathPara>'
           f'</a14:m><a:endParaRPr lang="en-US" sz="{sz}"/></a:p></p:txBody></p:sp></mc:Choice>'
           f'<mc:Fallback><p:sp>{nv}{sppr}<p:txBody>{body}<a:p><a:pPr algn="l"/><a:r><a:rPr lang="en-US" sz="{sz}">'
           f'<a:solidFill><a:srgbClr val="1F3A5F"/></a:solidFill><a:latin typeface="Cambria Math"/><a:ea typeface="微软雅黑"/></a:rPr>'
           f'<a:t>{fb}</a:t></a:r></a:p></p:txBody></p:sp></mc:Fallback></mc:AlternateContent>')
    xml = xml.replace("<m:", "<m:").replace("m:val", "m:val")
    # m: prefix inside runs needs the math namespace in scope -> declare on the AlternateContent root too
    xml = xml.replace('xmlns:p="', f'xmlns:m="{M_NS}" xmlns:p="', 1)
    el = etree.fromstring(xml)
    slide.shapes._spTree.append(el)
    return el


# ---------------- doughnut ----------------

def doughnut(slide, x, y, w, h, cats, vals, colors, hole=58, label_color=WHITE, label_size=14):
    cd = CategoryChartData(); cd.categories = cats; cd.add_series("条数", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = "章节环形图(原生)"
    ch = gf.chart
    ch.has_legend = False; ch.has_title = False
    plot = ch.plots[0]
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid(); pt.format.fill.fore_color.rgb = c
        pt.format.line.color.rgb = WHITE; pt.format.line.width = Pt(2)
    dl = plot.data_labels
    dl.show_value = True; dl.show_category_name = False
    dl.font.size = Pt(label_size); dl.font.bold = True; dl.font.color.rgb = label_color
    cs = ch._chartSpace
    dn = cs.find(".//" + qn("c:doughnutChart"))
    hs = dn.find(qn("c:holeSize"))
    if hs is None:
        hs = etree.SubElement(dn, qn("c:holeSize"))
    hs.set("val", str(hole))
    for d in cs.iter(qn("a:defRPr")):
        _set_fonts(d)
    return gf


def column_with_line(slide, x, y, w, h, cats, vals, line_val, vmin, vmax, major, bar_colors,
                     value_fmt='0.00', axis_fmt='0.0', line_label='"X̄ = "0.00', plot_box=(0.08, 0.06, 0.72, 0.8)):
    """Columns + dashed mean line (data-bound, padding categories) on the same axes."""
    from charts import _el, _font_txpr
    cats2 = [" "] + list(cats) + ["  "]
    v2 = [None] + list(vals) + [None]
    cd = CategoryChartData(); cd.categories = cats2
    cd.add_series("结果", v2); cd.add_series("均值", [line_val] * len(cats2))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = "均值法示意图(原生)"
    ch = gf.chart; ch.has_legend = False; ch.has_title = False
    plot = ch.plots[0]; plot.gap_width = 55
    bar = plot.series[0]
    bar.format.fill.solid(); bar.format.fill.fore_color.rgb = TEAL
    for i, c in enumerate(bar_colors, start=1):
        pt = bar.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = c
    dl = bar.data_labels; dl.show_value = True; dl.number_format = value_fmt; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.INSIDE_BASE; dl.font.size = Pt(14); dl.font.bold = True; dl.font.color.rgb = WHITE
    va = ch.value_axis; va.minimum_scale = vmin; va.maximum_scale = vmax; va.major_unit = major
    va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = RULE
    va.tick_labels.number_format = axis_fmt; va.tick_labels.number_format_is_linked = False
    va.format.line.fill.background()
    ch.category_axis.format.line.color.rgb = RULE
    cs = ch._chartSpace
    plotArea = cs.find(".//" + qn("c:plotArea"))
    layout = plotArea.find(qn("c:layout"))
    if layout is None:
        layout = etree.Element(qn("c:layout")); plotArea.insert(0, layout)
    for c_ in list(layout):
        layout.remove(c_)
    ml = etree.SubElement(layout, qn("c:manualLayout"))
    for tag, v in (("c:layoutTarget", "inner"), ("c:xMode", "edge"), ("c:yMode", "edge"),
                   ("c:x", plot_box[0]), ("c:y", plot_box[1]), ("c:w", plot_box[2]), ("c:h", plot_box[3])):
        ml.append(_el(tag, v))
    barChart = plotArea.find(qn("c:barChart"))
    ax_ids = [e.get("val") for e in barChart.findall(qn("c:axId"))]
    ls = barChart.findall(qn("c:ser"))[1]; barChart.remove(ls)
    for t in ("c:invertIfNegative", "c:dLbls", "c:spPr", "c:dPt"):
        for e in ls.findall(qn(t)):
            ls.remove(e)
    spPr = etree.Element(qn("c:spPr")); ln = etree.SubElement(spPr, qn("a:ln")); ln.set("w", "25400")
    sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr")).set("val", str(ORANGE))
    etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    ls.find(qn("c:tx")).addnext(spPr)
    mk = etree.Element(qn("c:marker")); mk.append(_el("c:symbol", "none")); spPr.addnext(mk)
    dLbls = etree.Element(qn("c:dLbls")); d = etree.SubElement(dLbls, qn("c:dLbl")); d.append(_el("c:idx", len(cats2) - 1))
    nf = etree.SubElement(d, qn("c:numFmt")); nf.set("formatCode", line_label); nf.set("sourceLinked", "0")
    sp = etree.SubElement(d, qn("c:spPr")); etree.SubElement(sp, qn("a:noFill")); l2 = etree.SubElement(sp, qn("a:ln")); etree.SubElement(l2, qn("a:noFill"))
    d.append(_font_txpr("c:txPr", 14, ORANGE, bold=True))
    for tag, v in (("c:dLblPos", "r"), ("c:showLegendKey", 0), ("c:showVal", 1), ("c:showCatName", 0),
                   ("c:showSerName", 0), ("c:showPercent", 0), ("c:showBubbleSize", 0)):
        d.append(_el(tag, v))
    for tag, v in (("c:showLegendKey", 0), ("c:showVal", 0), ("c:showCatName", 0), ("c:showSerName", 0),
                   ("c:showPercent", 0), ("c:showBubbleSize", 0)):
        dLbls.append(_el(tag, v))
    mk.addnext(dLbls); ls.append(_el("c:smooth", 0))
    lc = etree.Element(qn("c:lineChart")); lc.append(_el("c:grouping", "standard")); lc.append(_el("c:varyColors", 0))
    lc.append(ls); lc.append(_el("c:marker", 1))
    for a in ax_ids:
        lc.append(_el("c:axId", a))
    barChart.addnext(lc)
    for ax_tag, size in (("c:catAx", 15), ("c:valAx", 13)):
        ax = plotArea.find(qn(ax_tag)); old = ax.find(qn("c:txPr"))
        if old is not None:
            ax.remove(old)
        ax.find(qn("c:crossAx")).addprevious(_font_txpr("c:txPr", size, MUTED if ax_tag == "c:valAx" else INK))
        mt = ax.find(qn("c:majorTickMark"))
        if mt is not None:
            mt.set("val", "none")
    for dd in cs.iter(qn("a:defRPr")):
        if dd.find(qn("a:ea")) is None:
            _set_fonts(dd)
    return gf
