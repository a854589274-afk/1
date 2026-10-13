"""Section ②: how to compare (pages 10–24)."""
import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from decklib import *
from decklib2 import *
from charts import bar_with_threshold

T2 = "② 怎么比"
SUBHDR = RGBColor(0xC9, 0xD4, 0xE0)
TEAL_TINT = RGBColor(0xDD, 0xF0, 0xEE)


def p10_compare_table(prs, s, ctx):
    hoist_from_group(s, 15, {14})
    clear_shapes(s, keep_ids={5, 14})
    gl_pic, ws_pic = shape_by_id(s, 5), shape_by_id(s, 14)
    chrome(s, T2, "两份文件对照：WS/T 407-2012 与 CNAS-GL047:2021", "WS/T 407 §1/§4–§6.9　｜　GL047 前言/§1–§8", badge_w=4.6)
    colw = [1.05, 5.54, 5.54]
    TX, TY = MX, 1.55
    rows = ctx["compare_rows"]
    gf, tbl = add_table(s, len(rows) + 1, 3, TX, TY, sum(colw), 0.7 + 0.4 * len(rows), colw, name="对照表(原生)")
    tbl.rows[0].height = Inches(0.7)
    for r in range(1, len(rows) + 1):
        tbl.rows[r].height = Inches(0.4)
    thumb_h = 0.6
    gl_w = thumb_h * 944 / (1392 * (1 - 0.06 - 0.42))
    ws_w = thumb_h * 1400 / 474
    hdr = [("维度", None, 0.1), ("WS/T 407-2012", "卫生行业标准", ws_w + 0.25), ("CNAS-GL047:2021", "CNAS认可指南", gl_w + 0.25)]
    for c, (main, sub, left) in enumerate(hdr):
        paras = [[(main, {"bold": True, "size": 18})]] if sub is None else [[(main, {"bold": True, "size": 18})], [(sub, {"size": 12, "color": SUBHDR})]]
        style_cell(tbl.cell(0, c), paras, color=WHITE, fill=NAVY, margins=(left, 0.03, 0.08, 0.03), line_pt=21)
        cell_borders(tbl.cell(0, c), None, WHITE if c < 2 else None, None, None, width=1.5)
    for r, (dim, ws, gl) in enumerate(rows, start=1):
        style_cell(tbl.cell(r, 0), [[(dim, {"bold": True})]], size=15, color=NAVY, fill=CARD)
        style_cell(tbl.cell(r, 1), ws, size=14, margins=(0.08, 0.03, 0.06, 0.03))
        style_cell(tbl.cell(r, 2), gl, size=14, margins=(0.08, 0.03, 0.06, 0.03))
        for c in range(3):
            cell_borders(tbl.cell(r, c), None, None, None, RULE, width=0.75)
    gl_pic.crop_top, gl_pic.crop_bottom = 0.06, 0.42
    place(ws_pic, TX + colw[0] + 0.12, TY + 0.05, ws_w, thumb_h)
    place(gl_pic, TX + colw[0] + colw[1] + 0.12, TY + 0.05, gl_w, thumb_h)
    ws_pic.name, gl_pic.name = "WS/T 407封面(原截图)", "GL047封面(原截图)"
    for p in (ws_pic, gl_pic):
        sp = s.shapes._spTree; sp.remove(p._element); sp.append(p._element)
    set_notes(s, ctx["compare_notes"])


def p11_gl047_divider(prs, s, ctx):
    clear_shapes(s)
    set_bg(s, NAVY)
    bg = s.shapes.add_picture(os.path.join(ctx["art"], "divider_bg_16x9.png"), 0, 0, prs.slide_width, prs.slide_height)
    bg.name = "分隔页背景(无文字)"; send_to_back(s, bg)
    tag = add_box(s, MX + 0.25, 1.75, 1.6, 0.4, fill=RGBColor(0x2B, 0x4A, 0x70), radius=0.16, name="章节标签")
    box_text(tag, [[(T2, {"bold": True})]], size=13, color=RGBColor(0x7F, 0xD3, 0xCD), margin=(0.12, 0, 0.08, 0))
    add_text(s, MX + 0.25, 2.4, 8.0, 1.0, [[("CNAS-GL047:2021", {"bold": True})]], size=48, color=WHITE, margin=(0, 0, 0, 0))
    add_text(s, MX + 0.25, 3.45, 7.5, 1.3, [[("医学实验室", {})], [("定量检验程序结果可比性验证指南", {})]], size=28, color=WHITE,
             margin=(0, 0, 0, 0), line_spacing=1.15)
    add_text(s, MX + 0.25, 5.0, 6.0, 0.5, [[("2021-04-25 发布实施", {})]], size=20, color=SUBHDR, margin=(0, 0, 0, 0))
    set_notes(s, "CNAS-GL047:2021，2021-04-25发布并实施，首次发布；规范性引用文件只有 WS/T 407-2012。CNAS《实验室认可规范文件清单》（2026-9-23）仍列为有效。前言：本文件为指南文件，医学实验室可参考使用，也可使用其它适宜方式（如应用临床数据进行统计评估）进行可比性验证。")


def p12_terms(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "术语：先把“比什么”说清楚", "GL047 §3 · WS/T 407 §3.1", badge_w=2.9)
    terms = [("分析物", "analyte", "具有可测量特性的样品组分"),
             ("被测量", "measurand", "拟测量的量：需说明量类、基质、化学实体"),
             ("测量程序", "measurement procedure", "对测量所作的详细描述，含必需的计算"),
             ("测量系统", "measuring system", "测量仪器及其他器械的组合，含试剂和用品")]
    cw, ch_, gx, gy = 3.55, 1.45, 0.2, 0.18
    for i, (cn, en, d) in enumerate(terms):
        x = MX + (i % 2) * (cw + gx); y = 1.62 + (i // 2) * (ch_ + gy)
        c = card(s, x, y, cw, ch_, name=f"术语卡{i+1}")
        box_text(c, [[(cn + "  ", {"bold": True, "size": 19, "color": NAVY}), (en, {"size": 12, "color": MUTED})],
                     [(d, {})]], size=15, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.18, 0.12, 0.12, 0.06))
    # decomposition
    y0 = 4.95
    add_text(s, MX, y0, 7.3, 0.4, [[("例：“24h尿蛋白质质量”——整个短语就是被测量", {"bold": True})]], size=16, color=NAVY, margin=(0, 0, 0, 0))
    parts = [("24h尿", "基质", NAVY), ("蛋白质", "分析物", TEAL), ("质量", "特性（量类）", ORANGE)]
    bw = 2.3
    for i, (w_, lab, col) in enumerate(parts):
        x = MX + i * (bw + 0.2)
        b = add_box(s, x, y0 + 0.5, bw, 0.62, fill=col, radius=0.12, name="拆解块")
        box_text(b, [[(w_, {"bold": True})]], size=20, color=WHITE, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
        add_text(s, x, y0 + 1.18, bw, 0.4, [[("＝ " + lab, {"bold": True})]], size=16, color=col if col != ORANGE else NAVY,
                 align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    # target image + comparability
    rx = 8.15; rw = SLIDE_W - MX - rx
    tp = picture_from_blob(s, ctx["target_blob"], rx, 1.62, rw, None, name="靶心图(原P7，已裁英文标注)")
    tp.crop_bottom = 0.21
    tp.height = Inches(rw * 439 * 0.79 / 1060)
    add_text(s, rx, 1.62 + rw * 439 * 0.79 / 1060 + 0.05, rw, 0.3, [[("准确度与精密度（示意，原P7）", {})]], size=12, color=MUTED, margin=(0, 0, 0, 0))
    c = card(s, rx, 3.65, rw, 3.0, fill=WHITE, line=TEAL, name="可比性定义")
    box_text(c, [[("可比性  ", {"bold": True, "size": 19, "color": NAVY}), ("comparability", {"size": 12, "color": MUTED})],
                 [("使用不同的测量程序测定某种分析物获得的检测结果间的一致性。结果间的差异不超过规定的可接受标准时，可认为结果具有可比性。", {})],
                 [("来源：WS/T 407 §3.1（GL047 未定义此术语）", {"size": 12, "color": TEAL, "bold": True})]],
             size=16, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.2, 0.15, 0.18, 0.1))
    set_notes(s, """【GL047 第3章原文】
分析物 analyte：具有可测量特性的样品组分。注：在“24h 尿蛋白质质量”中，“蛋白质”是分析物，“质量”是特性。在“血浆中葡萄糖物质浓度”中，“葡萄糖”是分析物，“浓度”是特性。两个例子中的整个短语代表被测量。[GB/T 29791.1-2013/ISO 18113-1:2009，定义3.3]
被测量 measurand：拟测量的量。注1：在检验医学中被测量的规定需说明量类（例如质量浓度）、含有该量的基质（例如血浆）以及涉及的化学实体（例如分析物）；注2：被测量可以是生物活性；注3：在化学上，“分析物”，或某种物质或化合物的名称，有时被用作“被测量”的术语。此用法是错误的，因为这些术语不指代量。[定义3.39]
测量程序 measurement procedure：按照一个或多个测量原理和给定的测量方法，基于一种测量模型，对测量所作的详细描述，包括获得测量结果所必需的任何计算。[定义3.44]
测量系统 measuring system：一套组合在一起的，适合给出规定类量规定区间内测得量值的一个或多个测量仪器以及经常和其他器械组成的组合，包括任何试剂和用品。[A3.37]
【WS/T 407 §3.1】可比性 comparability：使用不同的测量程序测定某种分析物获得的检测结果间的一致性。结果间的差异不超过规定的可接受标准时，可认为结果具有可比性。——原稿把此定义标在GL047下，已更正。
拆解图中“24h尿＝基质”依据被测量注1（基质例如血浆）推得。""")


def p13_metrological(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "计量学可比性：溯源到同一参照", "GL047 §3（GB/T 29791.1 A.3.39）· GL047 §8", badge_w=4.4)
    steps = ["SI单位", "一级参考测量程序／一级参考物质", "二级参考物质", "制造商工作校准品", "产品校准品", "患者样本结果"]
    for i, t in enumerate(steps):
        x = MX + i * 0.45; y = 1.68 + i * 0.78
        fill = NAVY if i == 0 else (TEAL if i == len(steps) - 1 else CARD)
        b = add_box(s, x, y, 3.95, 0.62, fill=fill, radius=0.1, name=f"溯源阶梯{i+1}")
        box_text(b, [[(t, {"bold": True})]], size=15, color=WHITE if fill != CARD else NAVY, margin=(0.15, 0, 0.08, 0))
    add_text(s, 4.85, 1.75, 1.9, 0.7, [[("↑ 计量溯源", {"bold": True})], [("↓ 量值传递", {"bold": True})]], size=13, color=MUTED,
             margin=(0, 0, 0, 0))
    add_text(s, MX, 6.42, 6.0, 0.3, [[("溯源链为示意（简化），具体以制造商溯源声明为准", {})]], size=12, color=MUTED, margin=(0, 0, 0, 0))
    rx = 7.0; rw = SLIDE_W - MX - rx
    c1 = card(s, rx, 1.68, rw, 1.75, fill=CARD, name="定义")
    box_text(c1, [[("测量结果的计量学可比性", {"bold": True, "color": NAVY})],
                  [("对于给定种类的量，计量可溯源到同一参照的测量结果间可比较的特性。", {})]],
             size=17, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.2, 0.14, 0.18, 0.1))
    c2 = card(s, rx, 3.65, rw, 1.55, fill=WHITE, line=TEAL, name="要点")
    box_text(c2, [[("溯源到同一参照", {"bold": True}), (" → ", {"color": TEAL, "bold": True}), ("具有计量学可比性", {"bold": True})],
                  [("→ 才可合理应用修正值/修正因子（见第8章）", {"bold": True, "color": TEAL, "size": 16})]],
             size=18, color=NAVY, anchor=MSO_ANCHOR.MIDDLE, margin=(0.2, 0.08, 0.15, 0.08))
    c3 = card(s, rx, 5.4, rw, 1.25, fill=CARD, name="反例")
    box_text(c3, [[("通常不具有计量学可比性：", {"bold": True})],
                  [("无有证标准物质、方法学原理不同、参考区间不同", {})]],
             size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE, margin=(0.2, 0.06, 0.18, 0.06))
    set_notes(s, """【GL047 第3章】测量结果的计量学可比性 metrological comparability of measurement results：对于给定种类的量，计量可溯源到同一参照的测量结果间可比较的特性。[GB/T 29791.1-2013/ISO 18113-1:2009，A.3.39]
【GL047 第8章】对于具有计量学可比性的检测系统结果，实验室可合理应用修正值或修正因子建立检测系统间结果的可比性，并保留相关记录。对于……不具有计量学可比性的检测系统（通常表现为没有可获得的有证标准物质、或检测方法学原理不同，或参考区间不同）……
原稿本页的截图定义的是“计量学溯源性”，与标题不符，已删除，改为原生溯源阶梯图（示意，简化，参考ISO 17511的一般结构）。""")


def _icon_card(s, x, y, w, h, ic, label, title_, sub, red=False, isz=1.15):
    c = card(s, x, y, w, h, fill=WHITE if red else CARD, line=RED if red else None, name="条件卡")
    if ic:
        icon(s, ic, x + 0.15, y + (h - isz) / 2, isz)
    tx = x + (isz + 0.3 if ic else 0.2)
    num_circle(s, tx, y + 0.18, 0.42, label, fill=RED if red else TEAL, size=15)
    add_text(s, tx + 0.52, y + 0.15, w - (tx - x) - 0.6, 0.5, [[(title_, {"bold": True})]], size=16, color=RED if red else NAVY,
             margin=(0, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, tx, y + 0.72, w - (tx - x) - 0.12, h - 0.8, [[(l, {"color": RED} if "【" in l else {})] for l in sub], size=14,
             color=INK, margin=(0, 0, 0, 0), line_spacing=1.1)


def p14_conditions(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "4.1 什么情况下要验证", "GL047 §4.1", badge_w=1.8)
    add_text(s, MX, 1.58, SLIDE_W - 2 * MX, 0.55, [[("同一分析物存在以下情况时，应验证不同检验程序在", {}), ("临床适宜区间", {"bold": True, "color": TEAL}),
                                                     ("内患者样品结果的可比性", {})]], size=17, color=INK, margin=(0, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE)
    items = [("c41_a_serum_plasma", "a", "样品类型不同", ["临床用途相同", "如血清与血浆葡萄糖", "（单位相同或可换算）"], False),
             ("c41_b_two_brands", "b", "不同检测系统", ["不同品牌/型号/方法"], False),
             ("c41_c_identical", "c", "多套相同系统", ["同型号多台"], False),
             ("c41_d_modules", "d", "同系统多模块", ["一台仪器的多个分析模块"], False),
             ("c41_e_sites", "e", "多地点/场所", ["中心、急诊、发热门诊", "【待填：我科实例，", "如门诊分部、POCT】"], False),
             (None, "×", "不宜比对", ["样品类型和临床用途都不同", "如动脉血气 vs 生化常规"], True)]
    w, h, gx, gy = 3.85, 2.12, 0.29, 0.2
    for i, (ic, lab, t, sub, red) in enumerate(items):
        x = MX + (i % 3) * (w + gx); y = 2.3 + (i // 3) * (h + gy)
        _icon_card(s, x, y, w, h, ic, lab, t, sub, red=red)
    set_notes(s, """【GL047 4.1 验证条件 原文】对于同一分析物（检验项目），实验室存在如下情况时，应验证不同检验程序在临床适宜区间（指对临床诊断、管理、预防、治疗或健康评估有意义的测量结果区间）内患者样品检验结果的可比性：
a) 检测的样品类型不同但临床预期用途相同，且测量单位相同或可换算时，如血清葡萄糖物质的量浓度与血浆葡萄糖物质的量浓度；注：检测的样品类型不同且临床预期用途不同时，不宜进行比对，如动脉血气检验项目与生化常规检验项目。
b) 使用不同的检测系统；c) 使用多套相同的检测系统；d) 使用同一检测系统的多个分析模块；e) 多地点或场所使用的检测系统，如中心实验室、急诊实验室、发热门诊实验室。
e)条下“我科实例”请科室填写。""")


def p15_timing(prs, s, ctx):
    pics = [sh.shape_id for sh in s.shapes]
    clear_shapes(s)
    chrome(s, T2, "4.2 什么时候验证", "GL047 §4.2 · WS/T 407 第5章", badge_w=3.0)
    stages = [("启用前", "全面验证", NAVY), ("常规运行", "定期评审", TEAL), ("要素变更", "评估影响，需要时重新验证", RGBColor(0x0E, 0x6E, 0x6A))]
    w = 3.95
    for i, (a, b, col) in enumerate(stages):
        x = MX + i * (w + 0.14)
        shp = s.shapes.add_shape(MSO_SHAPE.PENTAGON if i < 2 else MSO_SHAPE.PENTAGON, Inches(x), Inches(1.62), Inches(w), Inches(0.95))
        shp.name = f"时机{i+1}"; shp.fill.solid(); shp.fill.fore_color.rgb = col; shp.line.fill.background(); shp.shadow.inherit = False
        box_text(shp, [[(a + "：", {"bold": True}), (b, {})]], size=17, color=WHITE, margin=(0.2, 0, 0.35, 0))
    add_text(s, MX, 2.8, 9.0, 0.45, [[("WS/T 407 第5章：使用其比对方案的 8 种情形", {"bold": True})]], size=17, color=NAVY, margin=(0, 0, 0, 0))
    items = [("t42_qc_drift", "室内质控\n有漂移趋势"), ("t42_eqa_corrected", "室间质评不合格\n纠正后"),
             ("t42_reagent_lot", "更换试剂批号\n（必要时）"), ("t42_repair", "更换重要部件\n或重大维修后"),
             ("t42_software", "软件程序\n变更后"), ("t42_clinician", "临床医生对\n可比性有疑问"),
             ("t42_patient", "患者投诉有疑问\n（需确认时）"), ("t42_frequency", "需提高周期性\n比对频率")]
    w, h, gx, gy = 2.88, 1.62, 0.2, 0.18
    for i, (ic, t) in enumerate(items):
        x = MX + (i % 4) * (w + gx); y = 3.35 + (i // 4) * (h + gy)
        card(s, x, y, w, h, name=f"情形{i+1}")
        icon(s, ic, x + 0.1, y + (h - 0.9) / 2, 0.9)
        l1, l2 = t.split("\n")
        add_text(s, x + 1.08, y + 0.15, w - 1.12, h - 0.3, [[(f"{i+1} ", {"bold": True, "color": TEAL}), (l1, {})], [(l2, {})]], size=14,
                 color=INK, margin=(0, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)
    set_notes(s, """【GL047 4.2 验证时机 原文】
4.2.1 检测系统启用前，应进行全面的检测系统间可比性验证。
4.2.2 常规使用期间，实验室可利用日常工作产生的检验和质控数据，或临床医生反馈的意见，定期对检测系统间结果可比性进行评审，如不再满足检验结果预期用途的要求，应根据评估结果，采用适宜的方案，重新进行检测系统间可比性验证。
4.2.3 现用检测系统的任一要素（仪器、试剂、校准品等）变更，如仪器品牌或型号、试剂原理或成分、校准品溯源性等改变，应分析这些改变对检测系统间结果可比性的影响，需要时，采用适宜的方案，重新进行检测系统间可比性验证。
【WS/T 407 第5章 原文】5.1 室内质控结果有漂移趋势时。5.2 室间质评结果不合格，采取纠正措施后。5.3 更换试剂批号(必要时)。5.4 更换重要部件或重大维修后。5.5 软件程序变更后。5.6 临床医生对结果的可比性有疑问时。5.7 患者投诉对结果可比性有疑问(需要确认时)。5.8 需提高周期性比对频率时(如每季度或每月一次)。
另：WS/T 407 6.2 实施周期性结果比对时，用于检测临床样本的每个检测系统都应进行可比性验证。原稿截图已删除。""")


def p16_criteria(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "4.3 用什么标准判断", "GL047 §4.3 · WS/T 407 §6.7", badge_w=2.9)
    c = card(s, MX, 1.62, 4.45, 5.0, fill=CARD, name="GL047 4.3")
    box_text(c, [[("GL047 4.3", {"bold": True, "color": TEAL, "size": 14})],
                 [("实验室应根据检验项目的", {}), ("预期用途和性能要求", {"bold": True, "color": NAVY}), ("，制定适宜的检测系统结果间可比性的判断标准。", {})],
                 [(" ", {"size": 8})],
                 [("参考：制造商声明、国家/行业/地方/团体标准、临床指南、专家共识", {"size": 15, "color": MUTED})],
                 [(" ", {"size": 8})],
                 [("验证结果应满足实验室制定的判断标准", {"size": 15, "bold": True})],
                 [(" ", {"size": 8})],
                 [("WS/T 407 6.7.1 原则：", {"size": 14, "bold": True, "color": TEAL})],
                 [("满足临床需要，同时考虑检测系统的性能状况", {"size": 15})]],
             size=20, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.25, 0.2, 0.22, 0.1))
    rx = 5.35
    add_text(s, rx, 1.62, 7.4, 0.4, [[("WS/T 407 §6.7.2 优选顺序", {"bold": True}), ("（前提：各系统不精密度均符合要求）", {"size": 13, "color": MUTED})]],
             size=17, color=NAVY, margin=(0, 0, 0, 0))
    steps = ["a  临床研究推荐指标", "b  本院医生临床经验", "c  生物学变异：1/3 CVI", "d  室间质评（EQA）数据", "e  认可机构最低标准", "f  实验室长期不精密度"]
    for i, t in enumerate(steps):
        x = rx + i * 0.5; y = 2.15 + i * 0.6
        b = add_box(s, x, y, 4.35, 0.5, fill=NAVY if i == 0 else (TEAL if i == 2 else CARD), radius=0.1, name=f"优选{i+1}")
        box_text(b, [[(t, {"bold": True})]], size=16, color=WHITE if i in (0, 2) else NAVY, margin=(0.15, 0, 0.1, 0))
    g = add_box(s, rx, 5.95, SLIDE_W - MX - rx, 0.62, fill=WHITE, line=NAVY, radius=0.1, name="底线g")
    box_text(g, [[("g  底线：", {"bold": True}), ("所选分析质量要求至少应满足国家或行业标准", {})]], size=16, color=NAVY, margin=(0.15, 0, 0.1, 0))
    set_notes(s, """【GL047 4.3 判断标准 原文】实验室应根据检验项目的预期用途和性能要求，制定适宜的检测系统结果间可比性的判断标准。实验室制定判断标准时，应参考制造商或研发者声明的标准、国家标准、行业标准、地方标准、团体标准、公开发表的临床应用指南和专家共识等。验证结果应满足实验室制定的判断标准。
【WS/T 407 6.7.1】建立结果可接受标准的原则：可比性验证的可接受标准应满足临床需要，同时考虑检测系统的性能状况。如系统性能无法满足规定的比对标准，可比性验证将会经常失败，此时需要改进检测系统性能(更换检测系统或优化测量程序)以达到期望的比对标准。反之，如果基于检测系统不精密度建立的标准高于临床要求，实验室负责人可根据临床需要适当调整可接受标准。
【WS/T 407 6.7.2】确认参与比对的各检测系统的不精密度均符合要求的前提下，按以下优选顺序确定不同检测项目的分析质量要求：a) 依据临床研究结果得出的推荐指标；b) 依据医疗机构内医生的临床经验提出的建议指标；c) 依据生物学变异确定的分析质量要求（注：对于精密度检测结果符合要求的检测系统，同一检测项目可比性结果的允许差异为个体内生物学变异(CV)的1/3）；d) 依据室间质评(EQA)数据设定的分析质量要求；e) 依据认可机构设置的最低标准；f) 如无适用的外部标准，可依据实验室内部的长期不精密度数据确定分析质量要求；g) 所选定的分析质量要求至少应满足国家或行业标准的要求。
说明：g)是对所选要求的底线约束，不是第7个候选来源，故单独画成底线。原稿《全国临床检验操作规程》封面图已删除。""")


def p17_plan_prep(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "5–6 验证方案与实验前准备", "GL047 §5–§6 · WS/T 407 §6.3.1", badge_w=3.4)
    add_text(s, MX, 1.6, 4.0, 0.45, [[("第5条：程序文件/作业指导书应写明", {"bold": True})]], size=16, color=NAVY, margin=(0, 0, 0, 0))
    req = ["检测系统", "样本类型和数量", "实验操作", "质量控制过程", "数据处理方法", "判断标准"]
    for i, t in enumerate(req):
        y = 2.15 + i * 0.66
        num_circle(s, MX, y, 0.48, str(i + 1), fill=NAVY, size=15)
        b = add_box(s, MX + 0.6, y, 3.35, 0.48, fill=CARD, radius=0.1, name="要素")
        box_text(b, [[(t, {"bold": True})]], size=16, color=NAVY, margin=(0.15, 0, 0.1, 0))
    add_text(s, MX, 6.15, 4.2, 0.5, [[("负责人员：了解方案→制定计划→实施并记录", {})]], size=13, color=MUTED, margin=(0, 0, 0, 0))
    rx = 5.0; rw = SLIDE_W - MX - rx
    w3 = (rw - 2 * 0.18) / 3
    top = [("p6_technician", "人员", "熟悉方法原理与日常操作（样本处理、校准、维护、质控）"),
           ("p6_analyzer", "仪器", "关键性能指标经验证满足要求；辅助设备与标称值相符"),
           ("p6_reagents", "试剂与校准品", "验证过程中\n不宜更换批号")]
    for i, (ic, t, d) in enumerate(top):
        x = rx + i * (w3 + 0.18)
        card(s, x, 1.62, w3, 2.3, name=f"准备-{t}")
        icon(s, ic, x + 0.15, 1.77, 0.82)
        add_text(s, x + 1.05, 1.95, w3 - 1.1, 0.5, [[(t, {"bold": True})]], size=16, color=NAVY, margin=(0, 0, 0, 0))
        add_text(s, x + 0.2, 2.85, w3 - 0.35, 1.0, [[(l, {})] for l in d.split("\n")], size=14, color=INK, margin=(0, 0, 0, 0), line_spacing=1.1)
    sw = 2 * w3 + 0.18
    card(s, rx, 4.1, sw, 2.55, name="准备-样本")
    icon(s, "p6_tubes", rx + 0.15, 4.22, 0.95)
    add_text(s, rx + 1.2, 4.3, 2.5, 0.5, [[("样本", {"bold": True})]], size=17, color=NAVY, margin=(0, 0, 0, 0))
    half = (sw - 0.45) / 2
    a = card(s, rx + 0.15, 5.25, half, 1.28, fill=WHITE, line=TEAL, name="样本-GL047")
    box_text(a, [[("GL047 6.4：", {"bold": True, "color": TEAL}), ("应使用", {"bold": True}), ("患者样本（抗凝剂相同），≥5例，覆盖临床适宜区间，重点关注医学决定水平", {})]],
             size=13, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.12, 0.08, 0.1, 0.06))
    b = card(s, rx + 0.3 + half, 5.25, half, 1.28, fill=WHITE, line=NAVY, name="样本-WS407")
    box_text(b, [[("WS/T 407 6.3.1：", {"bold": True, "color": NAVY}), ("宜首选", {"bold": True}), ("临床标本；不得不使用其他物质时，应验证互通性", {})]],
             size=13, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.12, 0.08, 0.1, 0.06))
    x = rx + sw + 0.18
    card(s, x, 4.1, w3, 2.55, name="准备-质控品")
    icon(s, "p6_controls", x + 0.15, 4.22, 0.95)
    add_text(s, x + 1.2, 4.3, w3 - 1.3, 0.5, [[("质控品", {"bold": True})]], size=17, color=NAVY, margin=(0, 0, 0, 0))
    add_text(s, x + 0.2, 5.3, w3 - 0.35, 1.2, [[("使用适宜的质控品做好室内质量控制", {})]], size=14, color=INK, margin=(0, 0, 0, 0))
    set_notes(s, """【GL047 5 验证方案】实验室应制定验证程序文件和/或作业指导书，内容包括用于结果比对的检测系统（包括但不限于分析仪器、试剂和校准品），样本类型和数量，实验操作和质量控制过程，数据处理方法和判断标准等。负责实施结果可比性验证的人员应了解验证方案，制定验证计划，组织实施并记录。
【GL047 6 实验前准备】6.1 人员：实验操作人员应熟悉检测系统的方法原理与日常操作，包括样本处理、校准、维护程序、质量控制等，确保检测系统工作状态正常。6.2 仪器设备：所用检测系统的关键性能指标应经过验证满足性能要求，对测量结果有重要影响的辅助设备的性能指标应与标称值相符。6.3 试剂和校准品：验证过程中，试剂或校准品不宜更换批号。6.4 样本：应使用患者样本（若为抗凝样本，应使用相同的抗凝剂）；样本中被测物浓度、活性等应能覆盖临床适宜区间，重点关注医学决定水平。比对的样本数量应不少于5例。使用更多的样本数量，可以增加验证结果的可靠性。6.5 质控品：验证过程中应使用适宜的质控品做好室内质量控制。
【WS/T 407 6.3.1】宜使用临床标本作为首选比对物质；不得不使用其他物质(如室间质评物或其他参考物质)时，应验证比对物质的互通性。
易错：原稿把GL047的“应使用患者样本”写成“应首选”，并把WS/T 407“验证互通性”的要求混进GL047。两份文件要求不同，已分开写。原稿程序文件截图移至“程序文件对照”页。""")


def p18_method_tree(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "7.1 选哪种方法", "GL047 §7.1 / §7.2.3", badge_w=2.6)
    root = add_box(s, MX, 3.0, 2.9, 1.1, fill=NAVY, radius=0.12, name="根")
    box_text(root, [[("参与比对的", {"bold": True})], [("检测系统数量", {"bold": True})]], size=18, color=WHITE, align=PP_ALIGN.CENTER)
    up = add_box(s, 4.6, 1.75, 4.1, 1.0, fill=TEAL, radius=0.12, name="参比法")
    box_text(up, [[("可选：与参比系统比对的方法", {"bold": True})], [("见 7.2.1", {"size": 14})]], size=17, color=WHITE, margin=(0.2, 0.04, 0.1, 0.04))
    dn = add_box(s, 4.6, 4.35, 4.1, 1.0, fill=TEAL, radius=0.12, name="均值法")
    box_text(dn, [[("可选：均值法", {"bold": True})], [("见 7.2.2", {"size": 14})]], size=17, color=WHITE, margin=(0.2, 0.04, 0.1, 0.04))
    arrow(s, MX + 2.9, 3.35, 4.58, 2.25, color=NAVY, width=2.25)
    arrow(s, MX + 2.9, 3.75, 4.58, 4.85, color=NAVY, width=2.25)
    add_text(s, 3.45, 2.3, 1.1, 0.5, [[("≤ 4", {"bold": True})]], size=24, color=ORANGE, margin=(0, 0, 0, 0))
    add_text(s, 3.45, 4.3, 1.1, 0.5, [[("> 4", {"bold": True})]], size=24, color=ORANGE, margin=(0, 0, 0, 0))
    r1 = card(s, 9.1, 1.6, SLIDE_W - MX - 9.1, 1.3, fill=CARD, name="参比系统怎么定")
    box_text(r1, [[("参比系统怎么定：", {"bold": True, "color": NAVY})], [("综合性能确认/验证、IQC与EQA表现、不确定度评估", {})]],
             size=14, color=INK, margin=(0.15, 0.06, 0.12, 0.06))
    r2 = card(s, 9.1, 4.25, SLIDE_W - MX - 9.1, 1.2, fill=CARD, name="均值法要点")
    box_text(r2, [[("以全部系统结果的均值为参考值，计算相对极差 R", {})]], size=14, color=INK, margin=(0.15, 0.06, 0.12, 0.06))
    arrow(s, 8.72, 2.25, 9.08, 2.25, color=MUTED, width=1.5)
    arrow(s, 8.72, 4.85, 9.08, 4.85, color=MUTED, width=1.5)
    q = card(s, MX, 5.85, 5.7, 0.78, fill=WHITE, line=NAVY, name="7.2.3问题")
    box_text(q, [[("需要按不精密度定浓度范围和重复次数？", {"bold": True})]], size=16, color=NAVY, margin=(0.2, 0, 0.1, 0))
    a = card(s, 7.0, 5.85, SLIDE_W - MX - 7.0, 0.78, fill=TEAL_TINT, line=TEAL, name="7.2.3答案")
    box_text(a, [[("7.2.3：比对方案可参考 WS/T 407", {"bold": True})]], size=16, color=NAVY, margin=(0.2, 0, 0.1, 0))
    arrow(s, MX + 5.72, 6.24, 6.98, 6.24, color=TEAL, width=2.25)
    set_notes(s, """【GL047 7.1 验证方法和条件 原文】实验室可根据实际使用情况，选择使用与参比系统比对的方法或均值法进行可比性验证。
实验室使用的检测系统数量≤4 时，可以选用与参比系统比对的方法。实验室应根据检测系统分析性能的确认或验证结果、室内质控(IQC)和室间质评的表现、不确定度评估等情况，综合评估后，确定实验室内的参比系统。实验室使用的其它检测系统检测结果与参比系统的测量结果进行比对，计算每个检测系统与参比系统检测结果的偏差，并依此评价可比性验证结果。
实验室使用检测系统数量>4 时，可以选用均值法。以全部系统结果的均值为参考值，计算全部检测系统结果的极差，并依此评价可比性验证结果。
【7.2.3】适用时，可依据实验室所用测量系统的不精密度，确定比对样本的浓度范围和重复检测次数，比对方案的制定可参考 WS/T 407《医疗机构内定量检验结果的可比性验证指南》。
注意：原文是“可以选用”，不是“必须用”。""")


def p19_reference(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "7.2.1 参比系统比对法：怎么判通过", "GL047 §7.2.1", badge_w=2.0)
    add_text(s, MX, 1.6, SLIDE_W - 2 * MX, 0.5, [[("每份样本：各系统与参比系统结果的偏差，对照4.3判断标准", {})]], size=17, color=INK, margin=(0, 0, 0, 0))
    for i in range(5):
        ok = i != 3
        d = add_box(s, MX + 0.2 + i * 1.15, 2.35, 0.9, 0.9, fill=TEAL if ok else RED, shape=MSO_SHAPE.OVAL, radius=None, name="样本点")
        box_text(d, [[("✓" if ok else "×", {"bold": True})]], size=30, color=WHITE, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    add_text(s, 6.6, 2.3, 6.1, 1.0, [[("5", {"bold": True, "color": ORANGE}), (" 份中 ", {}), ("4", {"bold": True, "color": ORANGE}),
                                       (" 份符合　→　", {}), ("通过", {"bold": True, "color": TEAL})]], size=30, color=NAVY,
             margin=(0, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE)
    steps = [("必要时", "适当增加样本量", NAVY), ("90%以上符合", "→ 通过", TEAL), ("≥20份仍不符合", "→ 分析其他影响因素，采取措施", RGBColor(0x0E, 0x6E, 0x6A))]
    w = 3.95
    for i, (a, b, col) in enumerate(steps):
        x = MX + i * (w + 0.14)
        shp = s.shapes.add_shape(MSO_SHAPE.PENTAGON if i < 2 else MSO_SHAPE.RECTANGLE, Inches(x), Inches(4.1), Inches(w), Inches(1.5))
        shp.name = f"扩样步骤{i+1}"; shp.fill.solid(); shp.fill.fore_color.rgb = col; shp.line.fill.background(); shp.shadow.inherit = False
        box_text(shp, [[(a, {"bold": True, "size": 22})], [(b, {})]], size=16, color=WHITE, margin=(0.25, 0.05, 0.35, 0.05))
    add_text(s, MX, 5.95, SLIDE_W - 2 * MX, 0.6, [[("提示：判断标准由实验室按4.3预先制定；样本须符合6.4（患者样本、≥5例、覆盖医学决定水平）", {})]],
             size=14, color=MUTED, margin=(0, 0, 0, 0))
    set_notes(s, """【GL047 7.2.1 参比系统比对方案 原文】
a) 使用不同检测系统检测符合 6.4 要求的样本，按要求进行校准和室内质控，记录检测结果。
b) 分别计算不同检测系统结果与参比系统结果的偏差，并与 4.3 实验室的判断标准进行比较。
c) 5 份样本中有 4 份检验结果的偏差符合实验室制定的判断标准，即为可比性验证通过。
d) 必要时，可适当增加检测样本量，如果 90%以上的样本检测结果偏差符合实验室制定的判断标准，即为可比性验证通过。若比对样本量达到 20份或以上时，比对结果仍不符合判断标准，实验室应对其它影响结果可比性的因素进行分析并采取相应措施。""")


def p20_mean_method(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "7.2.2 均值法：看相对极差 R", "GL047 §7.2.2", badge_w=2.0)
    SZ = 2600
    e1 = (m_bar(m_r("X", SZ)) + m_r("=", SZ, italic=False) +
          m_frac(m_sub("X", "1", SZ) + m_r("+", SZ, False) + m_sub("X", "2", SZ) + m_r("+⋯+", SZ, False) + m_sub("X", "n", SZ), m_r("n", SZ)))
    e2 = (m_r("R", SZ) + m_r("=", SZ, italic=False) +
          m_frac(m_sub("X", "max", SZ) + m_r("−", SZ, False) + m_sub("X", "min", SZ), m_bar(m_r("X", SZ))) + m_r("×100%", SZ, False))
    card(s, MX, 1.62, 6.0, 1.35, name="公式1底")
    add_equation(s, MX + 0.25, 1.72, 5.15, 1.15, e1, "X̄ = (X₁ + X₂ + … + Xₙ) / n", size=26, name="公式(1)")
    add_text(s, MX + 5.45, 1.62, 0.5, 1.35, [[("(1)", {})]], size=14, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0, 0, 0))
    card(s, MX, 3.12, 6.0, 1.35, name="公式2底")
    add_equation(s, MX + 0.25, 3.22, 5.15, 1.15, e2, "R = (Xmax − Xmin) / X̄ × 100%", size=26, name="公式(2)")
    add_text(s, MX + 5.45, 3.12, 0.5, 1.35, [[("(2)", {})]], size=14, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0, 0, 0))
    add_text(s, MX, 4.55, 6.0, 0.35, [[("Xmax / Xmin：各检测系统结果中的最大值 / 最小值", {})]], size=13, color=MUTED, margin=(0, 0, 0, 0))
    ok = card(s, MX, 5.0, 6.0, 0.62, fill=TEAL, name="通过")
    box_text(ok, [[("R 符合判断标准　→　通过", {"bold": True})]], size=17, color=WHITE, margin=(0.2, 0, 0.1, 0))
    loop = s.shapes.add_shape(MSO_SHAPE.CIRCULAR_ARROW, Inches(MX), Inches(5.78), Inches(0.85), Inches(0.85))
    loop.name = "回环箭头"; loop.fill.solid(); loop.fill.fore_color.rgb = ORANGE; loop.line.fill.background(); loop.shadow.inherit = False
    add_text(s, MX + 1.0, 5.75, 5.0, 0.95, [[("不符合：差异最大的两个系统不可比", {"bold": True})], [("剔除偏差较大者后重算 R，直到剩余系统符合", {})]],
             size=15, color=NAVY, margin=(0, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE)
    vals = [5.05, 5.20, 4.85, 5.45, 5.10]
    mean = sum(vals) / 5
    cx, cy, cw, chh = 7.0, 1.95, SLIDE_W - MX - 7.0, 3.9
    column_with_line(s, cx, cy, cw, chh, ["A", "B", "C", "D", "E"], vals, round(mean, 2), 0, 6, 1,
                     [TEAL, TEAL, NAVY, NAVY, TEAL], value_fmt="0.00", axis_fmt="0", line_label='"X̄ = "0.00')
    tag = add_box(s, cx, 1.58, 1.5, 0.34, fill=WHITE, line=RED, radius=0.1, name="示意数据标签")
    box_text(tag, [[("示意数据", {"bold": True})]], size=13, color=RED, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    add_text(s, cx + 1.65, 1.58, cw - 1.65, 0.34, [[("5台仪器同一样本；深色＝Xmax(D)与Xmin(C)", {})]], size=12, color=MUTED,
             margin=(0, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE)
    R = (max(vals) - min(vals)) / mean * 100
    add_text(s, cx, 5.95, cw, 0.7, [[("R = (5.45 − 4.85) / 5.13 × 100% = ", {}), (f"{R:.1f}%", {"bold": True, "color": ORANGE})]],
             size=17, color=NAVY, margin=(0, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE)
    ctx["mean_demo"] = (vals, mean, R)
    set_notes(s, f"""【GL047 7.2.2 均值法比对方案 原文】
a) 使用不同检测系统检测符合 6.4 要求的样本，按要求进行校准和室内质控，记录检测结果。
b) 按下式计算所有检测系统结果的均值(X̄)：X̄ = (X1 + X2 + X3 + … Xn)/n ……(1)，X1、X2、…XN——不同检测系统的结果；
c) 按下式计算所有检测系统结果的相对极差（R）：R = [(Xmax − Xmin)/X̄] × 100% ……(2)，Xmax——检测系统结果中的最大值；Xmin——检测系统结果中的最小值；
d) 将相对极差（R）与 4.3 实验室的判断标准进行比较。
e) 若 R 值符合实验室制定的判断标准，即为可比性验证通过。
f) 若 R 值不符合实验室制定的标准，表明结果差异最大的两个检测系统结果可比性不符合要求，分析并剔除偏差较大的检测系统的结果，按 7.2.2 重新计算 R 值，直到剩余检测系统结果符合可比性要求。
【示意数据】A 5.05、B 5.20、C 4.85、D 5.45、E 5.10（无单位，仅示意）；X̄ = {mean:.3f}；R = 0.60/{mean:.3f} = {R:.2f}%（已用代码复算）。
【公式说明】两个公式为PowerPoint原生公式（可双击编辑）；在不支持Office公式的软件（如LibreOffice）中显示为同内容的文字版。""")


def p21_wst407_flow(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "WS/T 407 §6.4–6.8 比对试验实施", "WS/T 407 §6.4–§6.8", badge_w=2.5)
    steps = [("估计长期CV", "长期CVᵢ\n尽量用6个月数据"), ("比较CV差异", "CVmax / CVmin\n< 2 倍？"), ("定浓度范围", "m总 × (1 ± 20%)"),
             ("定重复次数", "CV合并 =\n√(ΣCVᵢ²/n)\n→ 查表A.1"), ("计算比对偏差", "R=(x̄max−x̄min)\n÷ x̄总 × 100%"), ("判定", "R ≤ 要求？")]
    cw, gap = 1.86, 0.19
    for i, (t, f) in enumerate(steps):
        x = MX + i * (cw + gap)
        num_circle(s, x + (cw - 0.55) / 2, 1.65, 0.55, "①②③④⑤⑥"[i], fill=ORANGE if i == 5 else TEAL, size=20)
        if i < 5:
            arrow(s, x + (cw + 0.55) / 2 + 0.05, 1.925, x + cw + gap + (cw - 0.55) / 2 - 0.05, 1.925, color=RULE, width=2)
        add_text(s, x, 2.32, cw, 0.5, [[(t, {"bold": True})]], size=17, color=NAVY, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
        b = card(s, x, 2.9, cw, 1.1, fill=CARD, name=f"公式徽章{i+1}")
        box_text(b, [[(l, {"bold": True})] for l in f.split("\n")], size=14, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 margin=(0.04, 0.04, 0.04, 0.04))
    x2 = MX + 1 * (cw + gap)
    c2 = card(s, MX, 4.3, 2 * cw + gap + 1.0, 1.2, fill=WHITE, line=NAVY, name="分支②")
    box_text(c2, [[("② 否", {"bold": True, "color": RED}), ("：按 CLSI EP9-A2 和(或) EP15-A2", {})], [("确认检测系统间的结果可比性", {"bold": True})]],
             size=15, color=INK, margin=(0.18, 0.06, 0.15, 0.06))
    arrow(s, x2 + cw / 2, 4.02, x2 + cw / 2, 4.28, color=NAVY, width=1.5)
    x6 = MX + 5 * (cw + gap)
    rx = MX + 2 * cw + gap + 1.25
    c6 = card(s, rx, 4.3, SLIDE_W - MX - rx, 1.55, fill=WHITE, line=TEAL, name="分支⑥")
    box_text(c6, [[("⑥ R ≤ 要求：", {"bold": True, "color": TEAL}), ("该浓度水平所有系统可比", {})],
                  [("⑥ R > 要求：", {"bold": True, "color": NAVY}), ("均值差异最大的两个系统分别与规范操作系统比较，剔除偏差较大者，对剩余系统重做6.8.4–6.8.5", {})]],
             size=15, color=INK, margin=(0.18, 0.06, 0.15, 0.06))
    arrow(s, x6 + cw / 2, 4.02, x6 + cw / 2, 4.28, color=TEAL, width=1.5)
    w_ = card(s, rx, 6.0, SLIDE_W - MX - rx, 0.65, fill=WHITE, line=RED, name="易错")
    box_text(w_, [[("易错　", {"bold": True, "color": RED}), ("是“大于”不是“大于或等于”；是“均值差异最大的两个系统”", {})]], size=14, color=INK, margin=(0.18, 0, 0.12, 0))
    set_notes(s, """【WS/T 407 原文】
6.4.1 使用日常工作中质控品的检测数据估计不精密度，尽可能使用累计6个月的检测数据计算长期的变异系数(CV)，以保证不精密度的估计结果具有代表性。
6.4.2 比较不同检测系统不精密度的大小，确定最大CV与最小CV间的差异是否小于2倍。如小于2倍，可使用本指南规定的比对方案；如大于2倍，则应参照CLSI EP9-A2和(或)EP15-A2确认检测系统间的结果可比性。
6.4.3 CV合并=[(CV1²+CV2²+…+CVn²)/n]^1/2（前提：各系统长期CV由基本相同的检测次数计算得出）。
6.5 以 m总×(1±20%) 作为比对样本的浓度选择范围，m总为所有检测系统质控品均值的总均值。
6.6 使用CV合并值查找附录A表A.1，确定每台仪器比对样本的重复检测次数。
6.8.2/6.8.3 未重复检测直接比较各系统结果；重复检测则先算各系统均值，再算总均值。
6.8.4 选取差异最大的两个检测系统的均值计算极差，并除以所有检测系统结果的总均值以得出比对偏差(R)。
6.8.5 a) 如比对偏差小于或等于分析质量要求，得出的结论是：在该次评估的样本浓度水平，所有检测系统的结果都具有可比性；b) 如比对偏差大于分析质量要求，得出的结论是：均值差异最大的两个检测系统间的可比性不符合要求。将两个检测系统的结果分别与规范操作检测系统的结果进行比较，剔除偏差较大的检测系统的结果后，对剩余检测系统的结果按6.8.4和6.8.5的程序进行可比性分析，直到剩余检测系统的比对偏差小于或等于分析质量要求。
6.8.6 对于不符合可比性要求的检测系统，应分析原因，必要时采取相应的纠正措施，其后再将该检测系统与规范操作检测系统（如使用配套试剂、用配套校准物定期进行仪器校准、仪器性能良好、规范地开展室内质量控制、参加室间质量评价成绩优良、检测程序规范、人员经过良好培训的检测系统）的结果进行比对，确认比对结果符合分析质量要求。
【已更正的原稿错误】原第7点“比对偏差大于或等于分析质量要求”→“大于”（6.8.5 b）；“均值最大的两个系统”→“均值差异最大的两个检测系统”；第1点(2)末句原文为“确认检测系统间的结果可比性”，不是“确认检测系统的性能符合要求”。""")


def p22_alt_inputs(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "ALT示例①：已知条件与方案", "WS/T 407 附录B.1 表B.1/B.2", badge_w=3.1)
    colw = [1.4, 1.2, 1.2, 1.2, 1.2]
    gf, t = add_table(s, 5, 5, MX, 1.6, sum(colw), 2.2, colw, name="表B.1不精密度(原生)")
    t.cell(0, 0).merge(t.cell(1, 0)); t.cell(0, 1).merge(t.cell(0, 2)); t.cell(0, 3).merge(t.cell(0, 4))
    style_cell(t.cell(0, 0), [[("分析仪", {"bold": True})]], size=15, color=WHITE, fill=NAVY, align=PP_ALIGN.CENTER)
    style_cell(t.cell(0, 1), [[("质控水平1", {"bold": True})]], size=15, color=WHITE, fill=NAVY, align=PP_ALIGN.CENTER)
    style_cell(t.cell(0, 3), [[("质控水平2", {"bold": True})]], size=15, color=WHITE, fill=NAVY, align=PP_ALIGN.CENTER)
    for c, lab in zip(range(1, 5), ["均值 U/L", "CV %", "均值 U/L", "CV %"]):
        style_cell(t.cell(1, c), [[(lab, {"bold": True})]], size=13, color=NAVY, fill=CARD, align=PP_ALIGN.CENTER)
    data = [("A", "48.7", "2.35", "124.2", "2.14"), ("B", "46.5", "2.43", "117.4", "2.06"), ("总结果", "47.6", "2.39", "120.8", "2.10")]
    for r, row in enumerate(data, start=2):
        for c, v in enumerate(row):
            last = r == 4
            style_cell(t.cell(r, c), [[(v, {"bold": last or c == 0})]], size=15, color=NAVY if (last or c == 0) else INK,
                       fill=CARD if last else WHITE, align=PP_ALIGN.CENTER)
    for r in range(5):
        t.rows[r].height = Inches(0.42)
        for c in range(5):
            cell_borders(t.cell(r, c), None, None, None, RULE if r >= 1 else None)
    add_text(s, MX, 3.95, 6.2, 0.4, [[("表B.2（2个检测系统部分）：临界值 %", {"bold": True})]], size=15, color=NAVY, margin=(0, 0, 0, 0))
    colw2 = [1.4, 1.2, 1.2, 1.2, 1.2]
    gf2, t2 = add_table(s, 5, 5, MX, 4.38, sum(colw2), 1.95, colw2, name="表B.2临界值(原生)")
    hdr = ["检测次数", "CV合并 1%", "2%", "3%", "4%"]
    rows = [("2", "4.30", "8.60", "12.90", "17.20"), ("3", "2.67", "4.53", "6.80", "9.07"),
            ("4", "1.73", "3.46", "5.19", "6.92"), ("5", "1.46", "2.92", "4.38", "5.83")]
    for c, v in enumerate(hdr):
        style_cell(t2.cell(0, c), [[(v, {"bold": True})]], size=13, color=WHITE, fill=NAVY, align=PP_ALIGN.CENTER)
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row):
            hi = (r == 2 and c in (2, 3))
            style_cell(t2.cell(r, c), [[(v, {"bold": hi or c == 0})]], size=14, color=NAVY if (hi or c == 0) else INK,
                       fill=TEAL_TINT if (r == 2) else WHITE, align=PP_ALIGN.CENTER)
    for r in range(5):
        t2.rows[r].height = Inches(0.38)
        for c in range(5):
            cell_borders(t2.cell(r, c), None, None, None, RULE if r >= 1 else None)
    add_text(s, MX, 6.35, 6.2, 0.35, [[("CV合并 2.39% 介于2%–3%；6% 落在 4.53%–6.80% 之间 → 重复3次", {})]], size=13, color=MUTED, margin=(0, 0, 0, 0))
    rx = 7.25; rw = SLIDE_W - MX - rx
    tiles = [("总均值 m总", [("47.6", True), (" / ", False), ("120.8", True), (" U/L", False)]),
             ("CV合并", [("2.39%", True), (" / ", False), ("2.10%", True)]),
             ("样本浓度范围 m总×(1±20%)", [("38.1–57.1", True), (" / ", False), ("96.6–145.0", True), (" U/L", False)]),
             ("判定标准", [("6%", True), ("　（CVI 18.0% 的 1/3，标准示例取值）", False)]),
             ("每份样本重复检测", [("n = 3", True), (" 次", False)])]
    th, tg = 0.92, 0.1
    for i, (lab, parts) in enumerate(tiles):
        y = 1.6 + i * (th + tg)
        c = card(s, rx, y, rw, th, fill=CARD, name=f"结果磁贴{i+1}")
        runs = [(txt, {"bold": True, "color": ORANGE, "size": 24} if big else {"size": 14, "color": INK}) for txt, big in parts]
        box_text(c, [[(lab, {"bold": True, "size": 13, "color": MUTED})], runs], size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE,
                 margin=(0.2, 0.04, 0.12, 0.04))
    set_notes(s, """【WS/T 407 附录B.1 原文】
B.1.1 依据长期室内质控结果估计检测结果的不精密度，结果见表B.1。质控品浓度水平1：m总=(48.7+46.5)/2=47.6 U/L；质控品浓度水平2：m总=(124.2+117.4)/2=120.8 U/L。
B.1.2 以每个质控品浓度水平的总均值m总×(1±20%)作为样本浓度的选择范围。样本1：47.6 U/L×(1±20%)(38.1 U/L～57.1 U/L)；样本2：120.8 U/L×(1±20%)(96.6 U/L～145.0 U/L)。实际比对样本浓度：仪器A检测值分别为45.4 U/L和126.7 U/L的样本被选为比对试验样本。
B.1.3 根据6.7的方法确定比对试验的结果可接受标准，查阅资料未发现相关临床研究结果提供建议，也没有相关临床医生建议，而ALT的个体内生物学变异(CVI)为18.0%，故实验室将CVI的1/3，即6%作为分析质量要求。
B.1.4 样本1：在表A.1中找到2个检测系统比对的对应部分(见表B.2)，CV合并值2.39%位于2%～3%之间，而之前确定的分析质量要求为6%，位于重复3次的范围区间内(4.53%<6%<6.80%)，故选择重复检测3次。样本2：依照上述方法，确定样本2的重复检测次数为3次。
【复算（代码，全部正确）】CV合并 = √((2.35²+2.43²)/2) = 2.390%；√((2.14²+2.06²)/2) = 2.100%；范围 38.08–57.12、96.64–144.96 U/L；按2%–3%列线性插值，n=3 时临界值约 5.42% < 6%。
【关于2.67】表B.2中“检测3次、CV合并1%”一格为标准原值2.67，与同行2%列4.53、3%列6.80（折合每1%均为2.27）不一致；按 q0.95(k,k(n−1))/√n×CV合并 复算应为2.27。本例使用2%、3%两列，不受影响。页面保留原值。
【参考】现行 WS/T 403—2024 表A.1：ALT 允许不精密度6.0%、允许偏倚5.0%、允许总误差 6 U/L(≤40 U/L)/15.0%(>40 U/L)。6% 是标准示例取值，本科实际判定标准由科室按6.7.2确定。""")


def p23_alt_results(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "ALT示例②：比对结果与判定", "WS/T 407 附录B.1 表B.3/B.4", badge_w=3.2)
    colw3 = [1.2, 1.25, 1.25, 1.25, 1.25]
    gf3, t3 = add_table(s, 6, 5, MX, 1.6, sum(colw3), 2.7, colw3, name="表B.3样本检测结果(原生)")
    t3.cell(0, 0).merge(t3.cell(1, 0)); t3.cell(0, 1).merge(t3.cell(0, 2)); t3.cell(0, 3).merge(t3.cell(0, 4))
    style_cell(t3.cell(0, 0), [[("检测次数", {"bold": True})]], size=15, color=WHITE, fill=NAVY, align=PP_ALIGN.CENTER)
    style_cell(t3.cell(0, 1), [[("样本1 / (U/L)", {"bold": True})]], size=15, color=WHITE, fill=NAVY, align=PP_ALIGN.CENTER)
    style_cell(t3.cell(0, 3), [[("样本2 / (U/L)", {"bold": True})]], size=15, color=WHITE, fill=NAVY, align=PP_ALIGN.CENTER)
    for c, lab in zip(range(1, 5), ["分析仪A", "分析仪B", "分析仪A", "分析仪B"]):
        style_cell(t3.cell(1, c), [[(lab, {"bold": True})]], size=14, color=NAVY, fill=CARD, align=PP_ALIGN.CENTER)
    data3 = [("1", "45.4", "44.1", "126.7", "124.5"), ("2", "46.3", "45.9", "125.5", "122.7"),
             ("3", "46.7", "44.6", "126.2", "123.8"), ("均值", "46.1", "44.9", "126.1", "123.7")]
    for r, row in enumerate(data3, start=2):
        last = r == 5
        for c, v in enumerate(row):
            style_cell(t3.cell(r, c), [[(v, {"bold": last or c == 0})]], size=16, color=NAVY if (last or c == 0) else INK,
                       fill=CARD if last else WHITE, align=PP_ALIGN.CENTER)
    for r in range(6):
        t3.rows[r].height = Inches(0.42)
        for c in range(5):
            cell_borders(t3.cell(r, c), None, None, None, RULE if r >= 1 else None)
    colw4 = [2.0, 2.1, 2.1]
    gf4, t4 = add_table(s, 4, 3, MX, 4.6, sum(colw4), 1.9, colw4, name="表B.4数据分析(原生)")
    rows4 = [("统计结果", "样本1", "样本2"), ("总均值", "45.5 U/L", "124.9 U/L"), ("极差", "1.2 U/L", "2.4 U/L"), ("比对偏差 R", "2.64%", "1.92%")]
    for r, row in enumerate(rows4):
        for c, v in enumerate(row):
            if r == 0:
                style_cell(t4.cell(r, c), [[(v, {"bold": True})]], size=15, color=WHITE, fill=NAVY, align=PP_ALIGN.CENTER)
            elif r == 3 and c > 0:
                style_cell(t4.cell(r, c), [[(v, {"bold": True})]], size=20, color=ORANGE, fill=WHITE, align=PP_ALIGN.CENTER)
            else:
                style_cell(t4.cell(r, c), [[(v, {"bold": c == 0})]], size=16, color=NAVY if c == 0 else INK,
                           fill=CARD if c == 0 else WHITE, align=PP_ALIGN.CENTER)
            cell_borders(t4.cell(r, c), None, None, None, RULE if r >= 1 else None)
        t4.rows[r].height = Inches(0.46)
    CXo, CYo, CW, CH = 7.25, 2.15, 5.48, 3.75
    bar_with_threshold(s, CXo, CYo, CW, CH, ["样本1", "样本2"], (2.64, 1.92), threshold=6, vmax=7, plot_box=(0.11, 0.05, 0.87, 0.8))
    badge = add_box(s, CXo + 0.1, 1.6, 1.55, 0.44, fill=TEAL, radius=0.22, name="通过徽章")
    box_text(badge, [[("✓ 通过", {"bold": True})]], size=18, color=WHITE, align=PP_ALIGN.CENTER, margin=(0.05, 0, 0.05, 0))
    add_text(s, CXo, 6.0, CW, 0.75, [[("结论：两个浓度的R均≤6%，A、B两台生化分析仪ALT结果可比", {})]], size=18, color=INK,
             margin=(0.05, 0, 0, 0), name="结论")
    set_notes(s, ctx["altb_notes"])


def p24_fail(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T2, "第8章 不通过怎么办", "GL047 §8 · WS/T 407 §6.9", badge_w=2.9)
    c0 = card(s, MX, 1.6, 8.35, 0.62, fill=CARD, name="先定偏离系统")
    box_text(c0, [[("先确定偏离较大的系统：", {"bold": True}), ("依据4.3要求、EQA/PT、正确度验证、临床反馈", {})]], size=15, color=INK, margin=(0.18, 0, 0.1, 0))
    b1 = add_box(s, MX, 2.75, 2.3, 0.95, fill=NAVY, radius=0.12, name="不通过")
    box_text(b1, [[("可比性", {"bold": True})], [("不通过", {"bold": True})]], size=17, color=WHITE, align=PP_ALIGN.CENTER)
    dm = s.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(3.3), Inches(2.45), Inches(3.3), Inches(1.55))
    dm.name = "判断-计量学可比性"; dm.fill.solid(); dm.fill.fore_color.rgb = TEAL; dm.line.fill.background(); dm.shadow.inherit = False
    box_text(dm, [[("具有计量学", {"bold": True})], [("可比性？", {"bold": True})]], size=16, color=WHITE, align=PP_ALIGN.CENTER, margin=(0.1, 0, 0.1, 0))
    arrow(s, MX + 2.3, 3.22, 3.28, 3.22, color=NAVY, width=2)
    yes = card(s, 7.15, 2.6, SLIDE_W - MX - 7.15, 1.25, fill=WHITE, line=TEAL, name="是")
    box_text(yes, [[("是 → ", {"bold": True, "color": TEAL}), ("合理应用修正值或修正因子，", {})], [("建立系统间可比性，并保留记录", {})]], size=16, color=INK,
             margin=(0.18, 0.04, 0.12, 0.04))
    arrow(s, 6.62, 3.22, 7.13, 3.22, color=TEAL, width=2)
    no = card(s, MX, 4.3, 8.35, 0.82, fill=CARD, name="否")
    box_text(no, [[("否", {"bold": True, "color": RED}), ("：无有证标准物质、方法学原理不同、参考区间不同；", {})],
                  [("　或样本类型不同；或纠正/调整后仍不满足判断标准", {})]],
             size=14, color=INK, margin=(0.18, 0.02, 0.1, 0.02))
    arrow(s, 4.95, 4.02, 4.95, 4.28, color=NAVY, width=2)
    a = add_box(s, MX, 5.5, 3.3, 1.15, fill=NAVY, radius=0.12, name="避免发出")
    box_text(a, [[("避免向临床发出具有", {"bold": True})], [("不同临床意义或解释的结果", {"bold": True})]], size=15, color=WHITE, margin=(0.15, 0.04, 0.1, 0.04))
    add_text(s, MX + 3.35, 5.85, 0.5, 0.45, [[("或", {"bold": True})]], size=18, color=NAVY, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    b = add_box(s, MX + 3.9, 5.5, 3.3, 1.15, fill=NAVY, radius=0.12, name="与用户讨论")
    box_text(b, [[("与用户讨论不可比结果的", {"bold": True})], [("影响及如何正确应用", {"bold": True})]], size=15, color=WHITE, margin=(0.15, 0.04, 0.1, 0.04))
    rec = add_box(s, MX + 7.45, 5.5, 0.9, 1.15, fill=TEAL, radius=0.12, name="记录")
    box_text(rec, [[("记录", {"bold": True})]], size=16, color=WHITE, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    arrow(s, MX + 7.22, 6.08, MX + 7.43, 6.08, color=TEAL, width=2)
    arrow(s, 2.25, 5.14, 2.25, 5.48, color=NAVY, width=1.5); arrow(s, 6.15, 5.14, 6.15, 5.48, color=NAVY, width=1.5)
    side = card(s, 9.25, 4.3, SLIDE_W - MX - 9.25, 2.35, fill=WHITE, line=NAVY, name="WS407 6.9")
    box_text(side, [[("WS/T 407 §6.9 并入", {"bold": True, "color": NAVY})],
                    [("结果不可比且难以纠正时：", {})], [("与临床沟通，", {})],
                    [("采用不同的参考区间和(或)医学决定水平，并在检验报告单上明确标示", {"bold": True})]],
             size=14, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.16, 0.12, 0.12, 0.08))
    set_notes(s, """【GL047 8 相关措施 原文】
应根据 4.3 的要求、EQA/PT 结果、正确度验证结果及临床医生反馈的意见等，评估并确定结果偏离较大的检测系统。
对于具有计量学可比性的检测系统结果，实验室可合理应用修正值或修正因子建立检测系统间结果的可比性，并保留相关记录。
对于符合 4.2 规定的可比性验证条件，但不具有计量学可比性的检测系统（通常表现为没有可获得的有证标准物质、或检测方法学原理不同，或参考区间不同）、不同样本类型（如血清、血浆、全血）及经过纠正或调整，不同测量系统间的可比性仍不满足判断标准时，实验室应采取有效措施避免向临床发出具有不同临床意义或解释的结果，或与用户讨论不可比结果对临床活动的影响，以及如何正确应用不同测量系统的检测结果（包括结果偏倚情况），并记录。
注：原文“符合 4.2 规定的可比性验证条件”中的“4.2”，按GL047章节结构应为4.1（4.1为验证条件，4.2为验证时机），此处照录原文。
【WS/T 407 6.9】维持结果的可比性需以检测系统各质量保证环节的标准化为前提，必要时通过校准改善结果的可比性，即不同检测系统通过结果的数字转换获得结果的一致性；结果不可比且难以纠正时，应与临床进行沟通，采用不同的参考区间和(或)医学决定水平并在检验报告单上明确标示。""")
