"""Full deck: rebuild the ORIGINAL 31-slide pptx into the new 29-page structure.
Original slides are reused (and their screenshots kept) wherever a new page maps to one;
merged/obsolete slides are deleted; 5 new pages are added on the blank layout."""
import sys, os, re, zipfile, shutil
from pptx import Presentation
import decklib2
from decklib import *
from content_shared import COMPARE_ROWS, COMPARE_NOTES, ALTB_NOTES
import pages_a as A, pages_b as B, pages_c as Cc

SRC, OUT, ART = sys.argv[1], sys.argv[2], sys.argv[3]
decklib2.ICON_DIR = os.path.join(ART, "icons")
prs = Presentation(SRC)
S = {i + 1: s for i, s in enumerate(prs.slides)}
blank = prs.slide_layouts[6]
N = {k: prs.slides.add_slide(blank) for k in ("pyramid", "inventory", "procedure", "summary", "end")}

ctx = {"art": ART, "compare_rows": COMPARE_ROWS, "compare_notes": COMPARE_NOTES, "altb_notes": ALTB_NOTES,
       "lzjy_blob": shape_by_id(S[20], 4).image.blob}

plan = [  # (page function, slide)
    (A.p01_cover, S[1]), (A.p02_toc, S[2]), (A.p03_hr_measures, S[3]), (A.p04_timeline, S[4]),
    (A.p05_measures_overview, S[5]), (A.p06_quality_chain, S[6]), (A.p07_three_articles, S[7]),
    (A.p08_cl02, S[11]), (A.p09_pyramid, N["pyramid"]),
    (B.p10_compare_table, S[12]), (B.p11_gl047_divider, S[13]), (B.p12_terms, S[14]), (B.p13_metrological, S[16]),
    (B.p14_conditions, S[17]), (B.p15_timing, S[18]), (B.p16_criteria, S[19]), (B.p17_plan_prep, S[20]),
    (B.p18_method_tree, S[22]), (B.p19_reference, S[23]), (B.p20_mean_method, S[24]), (B.p21_wst407_flow, S[27]),
    (B.p22_alt_inputs, S[28]), (B.p23_alt_results, S[31]), (B.p24_fail, S[26]),
    (Cc.p25_inventory, N["inventory"]), (Cc.p26_procedure, N["procedure"]), (Cc.p27_summary, N["summary"]),
    (Cc.p28_end, N["end"]), (Cc.p29_backup, S[15]),
]
for fn, slide in plan:
    fn(prs, slide, ctx)

# ---- reorder + delete merged slides (orig 8, 9, 10, 21, 25, 29, 30)
order = [slide for _, slide in plan]
sldIdLst = prs.slides._sldIdLst
by_part = {}
for sldId in list(sldIdLst):
    by_part[prs.part.related_part(sldId.rId)] = sldId
keep_parts = {sl.part for sl in order}
for part, sldId in list(by_part.items()):
    sldIdLst.remove(sldId)
    if part not in keep_parts:
        prs.part.drop_rel(sldId.rId)
for sl in order:
    sldIdLst.append(by_part[sl.part])
prs.core_properties.title = "定量检验程序结果可比性验证"
prs.save(OUT)

# ---- theme fonts + any leftover legacy typefaces -> 微软雅黑 / Arial
LEGACY = ["宋体", "黑体", "仿宋", "仿宋_GB2312", "华文行楷", "楷体", "楷体_GB2312", "等线", "等线 Light", "system-ui",
          "-apple-system", "Calibri", "Calibri Light", "Times New Roman", "新宋体", "华文中宋", "方正小标宋简体"]
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.endswith(".xml") and item.filename.startswith("ppt/"):
            x = data.decode("utf-8")
            if re.match(r"ppt/theme/theme\d+\.xml", item.filename):
                def fix(m):
                    b = m.group(0)
                    b = re.sub(r'<a:latin typeface="[^"]*"', '<a:latin typeface="Arial"', b)
                    b = re.sub(r'<a:ea typeface="[^"]*"', '<a:ea typeface="微软雅黑"', b)
                    b = re.sub(r'<a:font script="Hans" typeface="[^"]*"', '<a:font script="Hans" typeface="微软雅黑"', b)
                    return b
                x = re.sub(r"<a:(majorFont|minorFont)>.*?</a:\1>", fix, x, flags=re.S)
            for f in LEGACY:
                x = re.sub(rf'<a:(ea|cs|sym) typeface="{re.escape(f)}"', r'<a:\1 typeface="微软雅黑"', x)
                x = re.sub(rf'<a:latin typeface="{re.escape(f)}"', '<a:latin typeface="Arial"', x)
            data = x.encode("utf-8")
        zout.writestr(item, data)
shutil.move(tmp, OUT)
print("saved", OUT, len(order), "slides")
