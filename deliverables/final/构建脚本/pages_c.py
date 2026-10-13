"""Section ③ (pages 25–27), end page, backup page."""
import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from decklib import *
from decklib2 import *

T3 = "③ 我科怎么做"
PENDING = RGBColor(0xB0, 0xB8, 0xC2)


def p25_inventory(prs, s, ctx):
    chrome(s, T3, "我科多系统项目清单（待科室填写）", "WS/T 407 §6.1–6.2 · GL047 §4.1/§7.1 · 《办法》第29条", badge_w=5.2)
    hdr = ["项目", "系统A", "系统B", "系统C…", "适用方案", "上次比对日期", "结果", "甘肃互认项目"]
    colw = [1.55, 1.45, 1.45, 1.45, 1.75, 1.6, 1.25, 1.63]
    items = ["生化", "血细胞分析", "凝血", "POCT血糖", "血气", "", ""]
    gf, t = add_table(s, len(items) + 1, len(hdr), MX, 1.6, sum(colw), 0.5 * (len(items) + 1), colw, name="多系统项目清单(原生)")
    for c, h in enumerate(hdr):
        style_cell(t.cell(0, c), [[(h, {"bold": True})]], size=14, color=WHITE, fill=NAVY, align=PP_ALIGN.CENTER)
    for r, it in enumerate(items, start=1):
        style_cell(t.cell(r, 0), [[(it if it else "【待填】", {"bold": bool(it)})]], size=15, color=NAVY if it else PENDING, fill=CARD)
        for c in range(1, len(hdr)):
            style_cell(t.cell(r, c), [[("【待填】", {})]], size=12, color=PENDING, fill=WHITE, align=PP_ALIGN.CENTER)
        for c in range(len(hdr)):
            cell_borders(t.cell(r, c), None, None, None, RULE)
    for r in range(len(items) + 1):
        t.rows[r].height = Inches(0.5)
    y = 1.6 + 0.5 * (len(items) + 1) + 0.25
    add_text(s, MX, y, SLIDE_W - 2 * MX, 0.4, [[("适用方案：", {"bold": True, "color": NAVY}), ("≤4个系统 → 参比系统比对法；>4个 → 均值法；需按CV定重复次数 → WS/T 407", {})]],
             size=15, color=INK, margin=(0, 0, 0, 0))
    add_text(s, MX, y + 0.45, SLIDE_W - 2 * MX, 0.4, [[("POCT血糖：", {"bold": True, "color": NAVY}), ("《办法》第29条要求床旁检验与实验室常规方法比对", {})]],
             size=15, color=INK, margin=(0, 0, 0, 0))
    set_notes(s, """【填写说明】
1. 每个在两个及以上检测系统上报告的定量项目填一行（WS/T 407 6.1：同一医疗机构内，使用多个检测系统向临床报告检验结果的检测项目均应进行可比性验证）。
2. “系统A/B/C…”填仪器品牌型号或编号；系统数决定方案（GL047 7.1：≤4个可选参比系统比对法，>4个可选均值法；WS/T 407 最多适用10个系统）。
3. “结果”填“通过/不通过/处理中”，不通过的项目按第8章处理并记录。
4. “甘肃互认项目”：军队医疗机构按属地规定参加互认（国卫医政发〔2024〕37号），请对照甘肃省现行互认项目清单【待核】填写Y/N。
5. 血气：GL047 4.1注——动脉血气与生化常规不宜比对；本行指血气分析仪之间的比对。""")


def p26_procedure(prs, s, ctx):
    chrome(s, T3, "程序文件对照：LZJY-2-CX-38《比对管理程序》", "GL047 §5 · A001:2023 §7.3.7.4 · LZJY-2-CX-38", badge_w=4.6)
    add_text(s, MX, 1.6, 4.1, 0.45, [[("要求写明（GL047 第5条）", {"bold": True})]], size=16, color=NAVY, margin=(0, 0, 0, 0))
    req = ["检测系统", "样本类型和数量", "实验操作", "质量控制过程", "数据处理方法", "判断标准"]
    for i, t in enumerate(req):
        y = 2.1 + i * 0.55
        num_circle(s, MX, y, 0.42, str(i + 1), fill=NAVY, size=14)
        add_text(s, MX + 0.55, y, 3.4, 0.42, [[(t, {})]], size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0, 0, 0))
    a = card(s, MX, 5.5, 4.0, 1.15, fill=WHITE, line=TEAL, name="A001补充")
    box_text(a, [[("A001:2023 7.3.7.4 另要求：", {"bold": True, "color": TEAL})], [("比对条件、频次、相关措施；记录授权审核签字、保存≥2年", {})]],
             size=13, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.14, 0.08, 0.1, 0.06))
    rx = 5.0
    cov = picture_from_blob(s, ctx["lzjy_blob"], rx, 1.6, None, 2.6, name="LZJY-2-CX-38首页(原P20截图)")
    add_text(s, rx + 2.05, 1.65, SLIDE_W - MX - rx - 2.05, 0.9, [[("LZJY-2-CX-38", {"bold": True, "size": 20})], [("《比对管理程序》", {"bold": True})]],
             size=18, color=NAVY, margin=(0, 0, 0, 0))
    add_text(s, rx + 2.05, 2.6, SLIDE_W - MX - rx - 2.05, 0.8, [[("第5/2版", {})], [("2022-08-30发布 · 2022-09-01实施", {})]],
             size=15, color=MUTED, margin=(0, 0, 0, 0))
    cw = (SLIDE_W - MX - rx - 0.36) / 3
    cards_ = [("已有", TEAL, WHITE, ["职责", "实施流程", "记录", "比对的监督", "相关文件"], None),
              ("程序文件中未见", NAVY, WHITE, ["样本类型与数量", "判断标准", "数据处理方法", "验证时机"], "可能在各专业组作业指导书中，请主任确认"),
              ("依据更新", RGBColor(0x2C, 0x5A, 0x86), WHITE, ["CNAS-CL02:2023", "CNAS-CL02-A001:2023", "WS/T 885—2026", "WS/T 406→2024版*"], "*如程序文件有引用"),]
    for i, (h, col, tc, lines, foot) in enumerate(cards_):
        x = rx + i * (cw + 0.18)
        hd = add_box(s, x, 4.45, cw, 0.5, fill=col, radius=0.1, name=f"对照-{h}")
        box_text(hd, [[(h, {"bold": True})]], size=15, color=tc, margin=(0.12, 0, 0.08, 0))
        body = card(s, x, 4.95, cw, 1.72, fill=CARD, radius=0.08, name=f"对照内容-{h}")
        paras = [[("▪ " + l, {})] for l in lines]
        if foot:
            paras.append([(foot, {"size": 11, "color": MUTED})])
        box_text(body, paras, size=13, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.12, 0.06, 0.08, 0.04))
    set_notes(s, """【用途】这是给主任决策用的对照清单，不是缺陷清单。
左栏：GL047 5 验证方案——程序文件和/或作业指导书内容包括用于结果比对的检测系统（包括但不限于分析仪器、试剂和校准品），样本类型和数量，实验操作和质量控制过程，数据处理方法和判断标准等。CNAS-CL02-A001:2023 7.3.7.4 1)另列出“比对条件、样品类型及数量、比对方案、判断标准、频次及相关措施”，3)比对记录应由授权人员审核并签字，并至少保留2年。
右栏：LZJY-2-CX-38《比对管理程序》（中国人民解放军联勤保障部队第九四〇医院检验科；版本号/修订次数5/2；2022年8月30日发布，2022年9月1日实施；共2页）。原稿截图可见章节：1目的、2范围、3职责、4工作程序（含外部比对流程、4.2.6比对结果离群原因）、4.3记录、4.4比对的监督、5相关文件。
“程序文件中未见”的四项是否写在各专业组作业指导书中，请主任确认。
依据更新：CNAS-CL02:2023（代替2012版，条款号改为7.3.7.4）；CNAS-CL02-A001:2023；WS/T 885—2026（4.2.3：实验室内部比对遵循WS/T 407等）；如程序文件引用WS/T 406-2012，该标准已废止，现行为WS/T 406—2024《临床血液检验常用项目分析质量标准》。""")


def p27_summary(prs, s, ctx):
    chrome(s, T3, "小结：三句话", "GL047 §4.2/§7/§8 · WS/T 407 §6.9", badge_w=3.6)
    items = [("t42_frequency", "什么时候比", [[("启用前", True), ("：全面验证", False)], [("常规运行", True), ("：定期评审", False)],
                                             [("要素变更", True), ("：评估后重新验证", False)]]),
             ("toc_tubes", "怎么比", [[("≤4个", True), ("：可选参比系统比对法", False)], [(">4个", True), ("：可选均值法", False)],
                                     [("按CV定重复次数", True), ("：WS/T 407", False)]]),
             ("toc_scales", "不过怎么办", [[("有计量学可比性", True), ("：", False)], [("修正值/修正因子，并记录", False)],
                                         [("否则", True), ("：与临床沟通，", False)], [("避免发出不同临床意义的结果", True)]])]
    w, gap, y, h = 3.85, 0.29, 1.75, 4.85
    for i, (ic, hd, runs) in enumerate(items):
        x = MX + i * (w + gap)
        card(s, x, y, w, h, name=f"小结卡{i+1}")
        icon(s, ic, x + 0.3, y + 0.3, 1.35)
        add_text(s, x + 1.8, y + 0.55, w - 1.9, 0.8, [[(hd, {"bold": True})]], size=24, color=NAVY, margin=(0, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, x + 0.3, y + 1.95, w - 0.5, 2.7, [[(t, {"bold": b, "color": TEAL if b else INK}) for t, b in line] for line in runs],
                 size=17, color=INK, margin=(0, 0, 0, 0), line_spacing=1.35)
    set_notes(s, "什么时候比：GL047 4.2.1–4.2.3（启用前全面验证；常规使用期间定期评审；任一要素变更时分析影响，需要时重新验证），WS/T 407第5章8种情形。\n怎么比：GL047 7.1（≤4个可以选用参比系统比对法；>4个可以选用均值法），7.2.3可参考WS/T 407。\n不过怎么办：GL047 第8章；WS/T 407 6.9（不可比且难以纠正时与临床沟通，采用不同参考区间和(或)医学决定水平并在报告单上明确标示）。")


def p28_end(prs, s, ctx):
    set_bg(s, NAVY)
    pic = s.shapes.add_picture(os.path.join(ctx["art"], "cover_16x9.png"), 0, 0, prs.slide_width, prs.slide_height)
    pic.name = "结束页插图(复用封面)"; send_to_back(s, pic)
    add_text(s, 0.85, 2.3, 6.0, 1.2, [[("谢谢", {"bold": True})]], size=60, color=WHITE, margin=(0, 0, 0, 0))
    add_text(s, 0.85, 3.6, 6.0, 0.6, [[("欢迎讨论", {})]], size=26, color=RGBColor(0xDD, 0xE6, 0xEF), margin=(0, 0, 0, 0))
    add_text(s, 0.85, 5.3, 5.0, 0.5, [[("第940医院检验科", {})]], size=16, color=RGBColor(0xA9, 0xB9, 0xCC), margin=(0, 0, 0, 0))
    set_notes(s, "结束页。插图复用封面。")


def p29_backup(prs, s, ctx):
    orig_text = "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    clear_shapes(s)
    chrome(s, "备份", "备份：测量区间与检出限（原P15）", "原P15（待修订）", badge_w=2.2)
    flag = add_box(s, SLIDE_W - MX - 3.0, 0.62, 3.0, 0.62, fill=RED, radius=0.12, name="待修订标注")
    box_text(flag, [[("此页内容待修订", {"bold": True})]], size=20, color=WHITE, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    c = card(s, MX, 1.6, 7.3, 5.05, fill=CARD, name="原文摘录")
    box_text(c, [[("原稿示例（照录，未修改）", {"bold": True, "color": NAVY, "size": 14})],
                 [("如化学发光法检测AFP的实验，LoB（空白检测限）为0.8ug/L，检出限LoD为2.6ug/L，LoQ（定量LoD）为3.3ug/L，使用该检测系统AFP报告结果时，若AFP检测结果低于0.8ug/L，则报告“分析物未检出，浓度小于0.8ug/L”；若AFP检测结果高于0.8ug/L小于2.6ug/L，则报告“有分析物检出，但不能定量，浓度小于2.6ug/L”；若AFP检测结果高于3.3ug/L，则直接报告检测结果。", {})],
                 [(" ", {"size": 6})],
                 [("测量区间：在规定条件下，可由给定测量仪器或测量系统以规定的仪器不确定度测量的相同类量的量值的集合。体外诊断医疗器械性能特征已被验证的测量区间被称为可报告范围；测量区间的下限不应和检出限混淆。（GL047 §3）", {"color": MUTED})]],
             size=14, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.2, 0.15, 0.18, 0.1))
    rx = 8.2; rw = SLIDE_W - MX - rx
    add_text(s, rx, 1.6, rw, 0.45, [[("待修订问题", {"bold": True})]], size=18, color=RED, margin=(0, 0, 0, 0))
    issues = ["2.6–3.3 μg/L 区间没有报告规则", "“未检出”的切点用的是 LoB", "“检出限……某元素”为分析化学教材措辞",
              "错字：“1各单位”应为“1个单位”", "EP17-A 已被 EP17-A2（2012）取代【待核】"]
    for i, t in enumerate(issues):
        y = 2.15 + i * 0.88
        num_circle(s, rx, y, 0.45, str(i + 1), fill=RED, size=15)
        add_text(s, rx + 0.6, y - 0.05, rw - 0.6, 0.8, [[(t, {})]], size=16, color=INK, margin=(0, 0, 0, 0), line_spacing=1.1)
    set_notes(s, "【原P15全文（照录，待修订）】\n" + orig_text + "\n\n【待修订问题】①2.6–3.3 μg/L区间没有报告规则；②“未检出”的切点用的是LoB；③“检出限…某元素”是分析化学教材的措辞；④错字“1各单位”应为“1个单位”；⑤EP17-A已被EP17-A2(2012)取代【待核】。")
