"""Section ①: why compare (pages 1–9)."""
import os
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from decklib import *
from decklib2 import *

T1 = "① 为什么要比"
LIGHT_ON_NAVY = RGBColor(0xDD, 0xE6, 0xEF)


def p01_cover(prs, s, ctx):
    clear_shapes(s)
    set_bg(s, NAVY)
    pic = s.shapes.add_picture(os.path.join(ctx["art"], "cover_16x9.png"), 0, 0, prs.slide_width, prs.slide_height)
    pic.name = "封面插图(无文字)"
    send_to_back(s, pic)
    add_text(s, 0.85, 1.55, 6.2, 1.9, [{"runs": [("定量检验程序", {})], "space_after": 0}, {"runs": [("结果可比性验证", {})]}],
             size=46, color=WHITE, bold=True, margin=(0, 0, 0, 0), name="主标题", line_spacing=1.05)
    add_text(s, 0.85, 3.65, 6.0, 0.95, [[("WS/T 407-2012 & CNAS-GL047:2021", {})], [("解读与我科应用", {})]],
             size=20, color=LIGHT_ON_NAVY, margin=(0, 0, 0, 0), name="副标题", line_spacing=1.15)
    add_text(s, 0.85, 5.25, 5.0, 1.2, [[("主讲人　", {"color": RGBColor(0xA9, 0xB9, 0xCC)}), ("哈小琴", {"bold": True})],
                                       [("第940医院检验科", {})],
                                       [("2026年　月　日", {"color": RGBColor(0xA9, 0xB9, 0xCC)})]],
             size=16, color=WHITE, margin=(0, 0, 0, 0), name="主讲信息(日期可编辑)", line_spacing=1.25)
    set_notes(s, "【封面】主标题沿用原稿。副标题：WS/T 407-2012 & CNAS-GL047:2021 解读与我科应用。\n"
                 "日期为可编辑文本（“2026年　月　日”），讲课前直接改。插图为代码绘制的扁平矢量图，无任何文字/标志。")


def p02_toc(prs, s, ctx):
    clear_shapes(s)
    chrome(s, "目录", "今天回答三个问题", "全片依据：WS/T 407-2012 · CNAS-GL047:2021", badge_w=4.2)
    items = [("①", "为什么检验结果\n必须可比？", "政策要求 · 认可要求 · 文件层级", "toc_scales"),
             ("②", "不同检验程序的\n结果怎么比？", "GL047 · WS/T 407 · ALT示例", "toc_tubes"),
             ("③", "我科怎么做？", "项目清单 · 程序文件对照 · 小结", "toc_building")]
    w, gap, y, h = 3.85, 0.29, 1.75, 4.85
    for i, (no, q, sub, ic) in enumerate(items):
        x = MX + i * (w + gap)
        card(s, x, y, w, h, name=f"问题卡{i+1}")
        icon(s, ic, x + (w - 1.9) / 2, y + 0.35, 1.9)
        add_text(s, x + 0.3, y + 2.45, w - 0.6, 0.6, [[(no, {"bold": True})]], size=30, color=TEAL, margin=(0, 0, 0, 0))
        add_text(s, x + 0.3, y + 3.05, w - 0.6, 1.0, [[(l, {"bold": True})] for l in q.split("\n")], size=22, color=NAVY,
                 margin=(0, 0, 0, 0), line_spacing=1.1)
        add_text(s, x + 0.3, y + 4.1, w - 0.6, 0.5, [[(sub, {})]], size=14, color=MUTED, margin=(0, 0, 0, 0))
    set_notes(s, "三个问题对应三部分：①政策与认可要求（为什么要比）；②CNAS-GL047:2021 与 WS/T 407-2012 的方法（怎么比）；③我科多系统项目清单、程序文件对照与小结（我科怎么做）。")


def p03_hr_measures(prs, s, ctx):
    pic = shape_by_id(s, 9)          # 医政医管局 notice screenshot
    clear_shapes(s, keep_ids={9})
    chrome(s, T1, "《医疗机构检查检验结果互认管理办法》", "国卫医发〔2022〕6号 第4/13/39条", badge_w=3.6)
    pic.crop_top, pic.crop_bottom, pic.crop_left, pic.crop_right = 0.40, 0.0, 0.0, 0.0
    pw = 4.2
    ph = pw * (1576 * 0.60) / 1778
    place(pic, MX, 1.62, pw, ph); pic.name = "互认办法通知截图(已裁导航栏)"
    s.shapes.add_picture(os.path.join(ctx["art"], "scene_mutual_16x9.png"), Inches(MX), Inches(1.62 + ph + 0.25),
                         Inches(pw), Inches(pw * 9 / 16)).name = "插图-互认(无文字)"
    facts = [("7章39条", "全文结构"), ("2022-03-01", "起施行（第三十九条）"), ("四部门", "联合印发，含中央军委后勤保障部卫生局")]
    x0, w, gap = 5.2, 2.36, 0.22
    for i, (big, small) in enumerate(facts):
        x = x0 + i * (w + gap)
        card(s, x, 1.62, w, 1.75, name=f"事实卡{i+1}")
        add_text(s, x + 0.18, 1.75, w - 0.3, 0.7, [[(big, {"bold": True})]], size=28 if i != 1 else 24, color=ORANGE,
                 margin=(0, 0, 0, 0), anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, x + 0.18, 2.48, w - 0.3, 0.8, [[(small, {})]], size=14, color=INK, margin=(0, 0, 0, 0), line_spacing=1.1)
    add_text(s, x0, 3.65, 7.5, 0.45, [[("第四条　互认原则：五个“以”", {"bold": True})]], size=18, color=NAVY, margin=(0, 0, 0, 0))
    prin = [("底线", "保障质量安全"), ("前提", "质量控制合格"), ("导向", "降低患者负担"), ("根本", "满足诊疗需求"), ("标准", "接诊医师判断")]
    for i, (k, v) in enumerate(prin):
        y = 4.22 + i * 0.5
        hi = (k == "前提")
        tag = add_box(s, x0, y, 1.05, 0.4, fill=TEAL if hi else NAVY, radius=0.12, name="原则标签")
        box_text(tag, [[(k, {"bold": True})]], size=15, color=WHITE, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
        add_text(s, x0 + 1.2, y, 6.0, 0.4, [[(v, {"bold": hi})]], size=17, color=TEAL if hi else INK,
                 anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0, 0, 0))
    set_notes(s, """【国卫医发〔2022〕6号 原文要点（中国政府网核对）】
发文机关：国家卫生健康委、国家中医药局、国家医保局、中央军委后勤保障部卫生局；成文日期2022年2月14日。
第二条：本办法所称检查结果，是指通过超声、X线、核磁共振成像、电生理、核医学等手段对人体进行检查，所得到的图像或数据信息；所称检验结果，是指对来自人体的材料进行生物学、微生物学、免疫学、化学、血液免疫学、血液学、生物物理学、细胞学等检验，所得到的数据信息。检查检验结果不包括医师出具的诊断结论。
第四条：医疗机构应当按照“以保障质量安全为底线，以质量控制合格为前提，以降低患者负担为导向，以满足诊疗需求为根本，以接诊医师判断为标准”的原则，开展检查检验结果互认工作。
第十三条：医疗机构检查检验结果互认标志统一为HR。……如：“全国HR”“京津冀HR”“北京市西城区HR”等。未按要求参加质量评价或质量评价不合格的检查检验项目，不得标注。
第十四条：……对于检验结果应当注明所使用的检测方法及参考区间。
第三十九条：本办法自2022年3月1日起施行。全文七章三十九条。
【讲解提示】“以质量控制合格为前提”——检验科能控制的就是这一条；院内多系统结果可比是质控合格的一部分。
左侧截图来自原稿（已裁去网站导航栏）；原稿右侧的条文截图属于“截图形式的文字”，已删除，条文见本备注。""")


def p04_timeline(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T1, "检查检验结果互认历程", "文号见备注 · WS/T 885—2026 §4.2.3", badge_w=3.6)
    nodes = [("2006", "原卫生部\n提出检查互认", "【文号待核】"),
             ("2009", "同级医疗机构\n检查结果互认", ""),
             ("2016", "京津冀首批\n27项·132家", ""),
             ("2022", "《互认办法》施行\nHR标识", ""),
             ("2024", "七部门\n指导意见", "2025年底地市域>200项\n2027年底省域>300项"),
             ("2026.11", "WS/T 885—2026\n实施", "4.2.3：内部比对\n遵循WS/T 407")]
    ly = 3.05
    xs = [MX + 0.95 + i * 2.05 for i in range(6)]
    arrow(s, MX + 0.2, ly, SLIDE_W - MX - 0.1, ly, color=RULE, width=3, head=True, name="时间轴")
    for i, (yr, txt, sub) in enumerate(nodes):
        last = i == len(nodes) - 1
        x = xs[i]
        dot = add_box(s, x - 0.18, ly - 0.18, 0.36, 0.36, fill=ORANGE if last else TEAL, shape=MSO_SHAPE.OVAL, radius=None, name="节点")
        add_text(s, x - 1.0, 2.15, 2.0, 0.6, [[(yr, {"bold": True})]], size=24, color=ORANGE if last else NAVY,
                 align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0), anchor=MSO_ANCHOR.BOTTOM)
        add_text(s, x - 0.98, 3.4, 1.96, 0.85, [[(l, {"bold": last})] for l in txt.split("\n")], size=16, color=INK,
                 align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0), line_spacing=1.1)
        if sub:
            col = RED if "待核" in sub else MUTED
            add_text(s, x - 1.0, 4.28, 2.0, 0.7, [[(l, {})] for l in sub.split("\n")], size=12, color=col,
                     align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    c = card(s, MX, 5.35, 7.6, 1.25, fill=CARD, name="旁白")
    box_text(c, [[("互认的技术前提", {"bold": True}), ("　＝　", {"color": TEAL, "bold": True}), ("院内结果可比", {"bold": True, "color": TEAL})]],
             size=24, color=NAVY, margin=(0.3, 0.05, 0.2, 0.05))
    c2 = card(s, 8.45, 5.35, SLIDE_W - MX - 8.45, 1.25, fill=WHITE, line=TEAL, name="我院提示")
    box_text(c2, [[("我院为军队医疗机构：", {"bold": True})], [("按属地（甘肃）规定参加互认", {})]], size=16, color=NAVY,
             margin=(0.2, 0.05, 0.15, 0.05))
    set_notes(s, """【各节点出处与核查层级】
2006：原卫生部办公厅《关于医疗机构间医学检验、医学影像检查互认有关问题的通知》（卫办医发〔2006〕32号）——【待核：标题与日期仅见二手资料】。
2009：《国务院关于印发医药卫生体制改革近期重点实施方案（2009—2011年）的通知》（国发〔2009〕12号），原文：“规范公立医院临床检查、诊断、治疗、使用药物和植（介）入类医疗器械行为，优先使用基本药物和适宜技术，实行同级医疗机构检查结果互认。”〔中国政府网原文已核〕
（同期：国办函〔2010〕67号；卫生部办公厅《关于加强医疗质量控制中心建设推进同级医疗机构检查结果互认工作的通知》，文号原稿写“卫办医改发〔2010〕108号”，地方文件引作“卫办医政发〔2010〕108号”【待核】；国发〔2012〕57号《卫生事业发展“十二五”规划》【原句待核】；国发〔2013〕40号《国务院关于促进健康服务业发展的若干意见》原文：“二级以上医疗机构检验对所有医疗机构开放，推动医疗机构间检查结果互认。”〔已核〕）
2016：京津冀三地卫计委2016-09-05联合宣布，首批27项临床检验结果、132家医疗机构（北京69、天津37、河北26），2016-10-01实施，报告单用“★”标识。〔媒体报道层已核：健康界2016-09-05；三地原始通知未核〕
2022：国卫医发〔2022〕6号《医疗机构检查检验结果互认管理办法》，2022-03-01施行；第十三条互认标志统一为HR。〔已核〕原稿引用的2021年征求意见稿已删除。
2024：国卫医政发〔2024〕37号《关于进一步推进医疗机构检查检验结果互认的指导意见》，七部门（含中央军委后勤保障部卫生局），成文2024-11-08。原文：“到2025年底……各地市域内医疗机构间互认项目超过200项。到2027年底，各省域内医疗机构间互认项目超过300项”；“军队有关医疗机构按照属地有关规定参加互认工作。”〔已核〕
2026：WS/T 885—2026《临床检验结果互认的基本技术条件及质量指标》，2026-05-25发布，2026-11-01实施。4.2.3：“按照相关国家或行业标准进行检验方法选择（确认）、性能验证、实验室内部比对，遵循 WS/T 229、WS/T 347、WS/T 406、WS/T 407、WS/T 408、WS/T 409、WS/T 505、WS/T 806、WS/T 807 的要求。”〔书目已核；4.2.3条为科室核对〕
甘肃省现行互认文件及项目清单【待核】。""")


def p05_measures_overview(prs, s, ctx):
    pic = shape_by_id(s, 2)          # 卫生部文件 screenshot
    clear_shapes(s, keep_ids={2})
    chrome(s, T1, "《医疗机构临床实验室管理办法》总览", "卫医发〔2006〕73号 · 国卫办医函〔2020〕560号", badge_w=4.4)
    ph = 4.95; pw = ph * 1421 / 1650
    place(pic, MX, 1.62, pw, ph); pic.name = "卫生部文件截图(原稿)"
    cats = ["总则", "一般规定", "质量管理", "安全管理", "监督管理", "附则"]
    vals = [5, 16, 11, 11, 9, 4]
    colors = [RGBColor(0xC9, 0xD5, 0xE1), RGBColor(0xB3, 0xC3, 0xD4), TEAL, RGBColor(0xC9, 0xD5, 0xE1), RGBColor(0xB3, 0xC3, 0xD4), RGBColor(0xC9, 0xD5, 0xE1)]
    dx, dy, dw = 5.05, 1.62, 3.6
    doughnut(s, dx, dy, dw, dw, cats, vals, colors, hole=60, label_color=NAVY, label_size=13)
    add_text(s, dx + dw / 2 - 0.9, dy + dw / 2 - 0.55, 1.8, 0.45, [[("六章", {"bold": True})]], size=18, color=NAVY,
             align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    add_text(s, dx + dw / 2 - 0.9, dy + dw / 2 - 0.12, 1.8, 0.6, [[("56", {"bold": True, "size": 34, "color": ORANGE}), ("条", {"bold": True})]],
             size=18, color=NAVY, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    names = ["第一章 总则", "第二章 一般规定", "第三章 实验室质量管理", "第四章 实验室安全管理", "第五章 监督管理", "第六章 附则"]
    lx = 8.95
    for i, (n, v) in enumerate(zip(names, vals)):
        y = 1.7 + i * 0.52
        hi = i == 2
        if hi:
            card(s, lx - 0.1, y - 0.04, SLIDE_W - MX - lx + 0.1, 0.48, fill=TEAL, radius=0.1, name="高亮第三章")
        add_text(s, lx, y, 2.9, 0.4, [[(n, {"bold": hi})]], size=15, color=WHITE if hi else INK, anchor=MSO_ANCHOR.MIDDLE, margin=(0.05, 0, 0, 0))
        add_text(s, SLIDE_W - MX - 0.75, y, 0.7, 0.4, [[(f"{v}条", {"bold": True})]], size=15, color=WHITE if hi else NAVY,
                 anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT, margin=(0, 0, 0.05, 0))
    add_text(s, 5.05, 5.4, 7.7, 0.4, [[("与本主题直接相关", {"bold": True})]], size=16, color=NAVY, margin=(0, 0, 0, 0))
    for i, t in enumerate(["第九条　实验室集中设置、统一管理、资源共享", "第二十四条　保证检测系统的完整性和有效性"]):
        add_text(s, 5.05, 5.85 + i * 0.42, 7.7, 0.4, [[("▪ ", {"color": TEAL}), (t, {})]], size=16, color=INK, margin=(0, 0, 0, 0))
    add_text(s, 5.05, 6.72, 7.7, 0.28, [[("2020年国卫办医函〔2020〕560号修订（仅第二十八条）", {})]], size=12, color=MUTED, margin=(0, 0, 0, 0))
    set_notes(s, """【卫医发〔2006〕73号】卫生部2006年2月27日印发，第五十六条：“本办法自2006年6月1日起施行。”（NCCL刊载全文核对）
六章条数：总则5（第1–5条）、医疗机构临床实验室管理的一般规定16（第6–21条）、医疗机构临床实验室质量管理11（第22–32条）、医疗机构临床实验室安全管理11（第33–43条）、监督管理9（第44–52条）、附则4（第53–56条）。
第九条：医疗机构临床实验室应当集中设置，统一管理，资源共享。
第二十四条：医疗机构临床实验室应当保证检测系统的完整性和有效性，对需要校准的检验仪器、检验项目和对临床检验结果有影响的辅助设备定期进行校准。
2020年7月国卫办医函〔2020〕560号仅修订第二十八条（删去“经卫生部认定的”）——科室已核；我方未能独立复核卫健委原件。
原稿“（强制）”标注已删除（本办法为部门规范性文件，以原文为准）。
原稿10条要点中只保留与本主题相关的第九条、第二十四条。
原稿“2000年1月31日卫生厅医政司批准立项起草……2006上半年正式发布”——【待核】，疑为“卫生部医政司”，出处不明，未上页面。""")


def p06_quality_chain(prs, s, ctx):
    clear_shapes(s)
    chrome(s, T1, "《办法》第三章：一条质量链", "卫医发〔2006〕73号 第22–32条 · 全国标准信息公共服务平台", badge_w=5.6)
    nodes = [("SOP", "第22条"), ("仪器·试剂\n耗材", "第23条"), ("校准与检测\n系统完整性", "第24条"), ("室内质控", "第25–27条"),
             ("室间质评", "第28–29条"), ("比对", "第29–30条"), ("质量记录", "第32条")]
    w, gap, y, h = 1.5, 0.27, 1.9, 1.35
    for i, (t, art) in enumerate(nodes):
        x = MX + i * (w + gap)
        hi = t == "比对"
        b = add_box(s, x, y, w, h, fill=ORANGE if hi else (NAVY if i < 3 else TEAL), radius=0.12, name=f"链节{i+1}")
        paras = [[(l, {"bold": True})] for l in t.split("\n")] + [[(art, {"size": 13})]]
        box_text(b, paras, size=17 if len(t) < 6 else 15, color=WHITE, align=PP_ALIGN.CENTER, margin=(0.04, 0.02, 0.04, 0.02))
        if i < len(nodes) - 1:
            arrow(s, x + w + 0.03, y + h / 2, x + w + gap - 0.03, y + h / 2, color=MUTED, width=2)
    n4x = MX + 3 * (w + gap); n5x = MX + 4 * (w + gap)
    arrow(s, n4x + w / 2, y + h + 0.05, n4x + w / 2, 3.65, color=MUTED, width=1.25)
    arrow(s, n5x + w / 2, y + h + 0.05, n5x + w / 2 + 0.9, 3.65, color=MUTED, width=1.25)
    card_text(s, 3.95, 3.7, 3.8, 1.3, [[("第27条", {"bold": True}), ("所引计划号＝GB/T 20468-2006", {})],
                                       [("2026-01-28已废止", {})], [("现行：WS/T 641—2026", {"bold": True, "color": TEAL})]],
              size=14, margin=(0.14, 0.08, 0.1, 0.06), name="注-第27条")
    card_text(s, 7.95, 3.7, 3.8, 1.3, [[("第31条", {"bold": True}), ("所引计划号＝GB/T 20470-2006", {})],
                                       [("2026-01-28已废止", {})]],
              size=14, margin=(0.14, 0.08, 0.1, 0.06), name="注-第31条")
    warn = card(s, MX, 5.35, SLIDE_W - 2 * MX, 1.3, fill=WHITE, line=RED, name="注意-非国家原文")
    box_text(warn, [[("注意　", {"bold": True, "color": RED}),
                     ("原稿P6–P10各条下的（一）（二）（三）款，以及“每年比对不少于2次”“比对样本不少于5个”“相对偏差不能大于1/2PT”，均不是国家原文（疑为某省实施细则，出处待核），已删除。", {})]],
             size=16, color=INK, margin=(0.2, 0.08, 0.2, 0.08))
    set_notes(s, """【国家原文（卫医发〔2006〕73号，NCCL刊载全文核对）】
第二十二条：医疗机构应当加强临床实验室质量控制和管理。医疗机构临床实验室应当制定并严格执行临床检验项目标准操作规程和检验仪器的标准操作、维护规程。
第二十三条：医疗机构临床实验室使用的仪器、试剂和耗材应当符合国家有关规定。
第二十四条：医疗机构临床实验室应当保证检测系统的完整性和有效性，对需要校准的检验仪器、检验项目和对临床检验结果有影响的辅助设备定期进行校准。
第二十五条：医疗机构临床实验室应当对开展的临床检验项目进行室内质量控制，绘制质量控制图。出现质量失控现象时，应当及时查找原因，采取纠正措施，并详细记录。
第二十六条：医疗机构临床实验室室内质量控制主要包括质控品的选择，质控品的数量，质控频度，质控方法，失控的判断规则，失控时原因分析及处理措施，质控数据管理要求等。
第二十七条：医疗机构临床实验室定量测定项目的室内质量控制标准按照《临床实验室定量测定室内质量控制指南》（GB/20032302-T-361）执行。——该计划号对应GB/T 20468-2006，2026-01-28复审废止；室内质控现行行业标准为WS/T 641—2026《临床检验定量测定室内质量控制标准》（2026-11-01实施，代替WS/T 641-2018）。
第二十八条（2020年修订后）：医疗机构临床实验室应当参加室间质量评价机构组织的临床检验室间质量评价。
第二十九条：医疗机构临床实验室参加室间质量评价应当按照常规临床检验方法与临床检验标本同时进行，不得另选检测系统，保证检验结果的真实性。医疗机构临床实验室对于室间质量评价不合格的项目，应当及时查找原因，采取纠正措施。医疗机构应当对床旁临床检验项目与临床实验室相同临床检验项目常规临床检验方法进行比对。
第三十条：医疗机构临床实验室应当将尚未开展室间质量评价的临床检验项目与其他临床实验室的同类项目进行比对，或者用其他方法验证其结果的可靠性。临床检验项目比对有困难时，医疗机构临床实验室应当对方法学进行评价，包括准确性、精密度、特异性、线性范围、稳定性、抗干扰性、参考范围等，并有质量保证措施。
第三十一条：医疗机构临床实验室室间质量评价标准按照《临床实验室室间质量评价要求》（GB/20032301-T-361）执行。——该计划号对应GB/T 20470-2006，2026-01-28复审废止（全国标准信息公共服务平台）。
第三十二条：医疗机构临床实验室应当建立质量管理记录，包括标本接收、标本储存、标本处理、仪器和试剂及耗材使用情况、校准、室内质控、室间质评、检验结果、报告发放等内容。质量管理记录保存期限至少为2年。
【原稿问题】原P6–P10所列条文含国家原文没有的款项；其中“全省三级以上医疗机构……省卫生厅临床检验中心……”“……相对偏差不能大于1/2PT可接受范围”的出处【待核】。原P8的L-J图、Z分数图、Youden图因分辨率低删除；原P9“临床检验中心”网络图删除。""")


def p07_three_articles(prs, s, ctx):
    target = shape_by_id(s, 8)
    ctx["target_blob"] = target.image.blob
    clear_shapes(s)
    chrome(s, T1, "与可比性直接相关的三条（国家原文）", "卫医发〔2006〕73号 第24/29/30条", badge_w=3.6)
    cards = [("第二十四条", [("医疗机构临床实验室应当保证检测系统的", {}), ("完整性和有效性", {"bold": True, "color": TEAL}),
                          ("，对需要校准的检验仪器、检验项目和对临床检验结果有影响的辅助设备定期进行校准。", {})], "检测系统"),
             ("第二十九条（末句）", [("医疗机构应当对", {}), ("床旁临床检验项目", {"bold": True, "color": TEAL}),
                                ("与临床实验室相同临床检验项目常规临床检验方法", {}), ("进行比对", {"bold": True, "color": TEAL}), ("。", {})], "床旁检验 / POCT"),
             ("第三十条", [("医疗机构临床实验室应当将", {}), ("尚未开展室间质量评价", {"bold": True, "color": TEAL}),
                        ("的临床检验项目与其他临床实验室的同类项目进行比对，或者用其他方法验证其结果的可靠性。", {})], "无EQA项目")]
    w, gap, y, h = 3.85, 0.29, 1.7, 4.95
    for i, (hd, runs, tag) in enumerate(cards):
        x = MX + i * (w + gap)
        card(s, x, y, w, h, name=f"条文卡{i+1}")
        hb = add_box(s, x + 0.25, y + 0.25, 2.4, 0.46, fill=NAVY, radius=0.12, name="条号")
        box_text(hb, [[(hd, {"bold": True})]], size=16, color=WHITE, margin=(0.12, 0, 0.08, 0))
        add_text(s, x + 0.25, y + 0.95, w - 0.5, 3.2, [runs], size=19, color=INK, margin=(0, 0, 0, 0), line_spacing=1.25)
        add_text(s, x + 0.25, y + h - 0.6, w - 0.5, 0.4, [[("→ " + tag, {"bold": True})]], size=15, color=TEAL, margin=(0, 0, 0, 0))
    set_notes(s, """以下均为国家原文（卫医发〔2006〕73号，NCCL刊载全文逐字核对）。
第二十四条：医疗机构临床实验室应当保证检测系统的完整性和有效性，对需要校准的检验仪器、检验项目和对临床检验结果有影响的辅助设备定期进行校准。
第二十九条（全文）：医疗机构临床实验室参加室间质量评价应当按照常规临床检验方法与临床检验标本同时进行，不得另选检测系统，保证检验结果的真实性。医疗机构临床实验室对于室间质量评价不合格的项目，应当及时查找原因，采取纠正措施。医疗机构应当对床旁临床检验项目与临床实验室相同临床检验项目常规临床检验方法进行比对。（原稿漏了末句）
第三十条（全文）：医疗机构临床实验室应当将尚未开展室间质量评价的临床检验项目与其他临床实验室的同类项目进行比对，或者用其他方法验证其结果的可靠性。临床检验项目比对有困难时，医疗机构临床实验室应当对方法学进行评价，包括准确性、精密度、特异性、线性范围、稳定性、抗干扰性、参考范围等，并有质量保证措施。
注意：国家原文没有“每年比对不少于2次”“比对样本不少于5个”“相对偏差不能大于1/2PT”。原P7的靶心图已移到“术语”页。""")


def p08_cl02(prs, s, ctx):
    pic = shape_by_id(s, 5)          # CL02:2012 cover (wmf)
    clear_shapes(s, keep_ids={5})
    chrome(s, T1, "认可准则：CNAS-CL02:2023 第7.3.7.4条", "CNAS-CL02:2023 前言 · CNAS-CL02-A001:2023 §7.3.7.4（第15页）", badge_w=5.7)
    ph = 4.7; pw = ph * 599 / 781
    place(pic, MX, 1.62, pw, ph); pic.name = "CL02封面(原稿截图-2012版)"
    lab = add_box(s, MX, 1.62 + ph - 0.55, pw, 0.55, fill=WHITE, line=RED, radius=0.06, name="待替换标签")
    box_text(lab, [[("【待替换为2023版封面】", {"bold": True})]], size=13, color=RED, align=PP_ALIGN.CENTER, margin=(0.02, 0, 0.02, 0))
    rx = MX + pw + 0.45; rw = SLIDE_W - MX - rx
    t1 = add_box(s, rx, 1.62, 2.2, 0.48, fill=CARD, radius=0.12, name="旧条款")
    box_text(t1, [[("2012版 5.6.4", {"bold": True})]], size=15, color=MUTED, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    arrow(s, rx + 2.3, 1.86, rx + 2.95, 1.86, color=TEAL, width=2.5)
    t2 = add_box(s, rx + 3.05, 1.62, 2.4, 0.48, fill=TEAL, radius=0.12, name="新条款")
    box_text(t2, [[("2023版 7.3.7.4", {"bold": True})]], size=15, color=WHITE, align=PP_ALIGN.CENTER, margin=(0, 0, 0, 0))
    add_text(s, rx, 2.22, rw, 0.4, [[("2023-06-01发布 · 2023-12-01实施 · 等同采用ISO 15189:2022 · 代替CL02:2012", {})]],
             size=13, color=MUTED, margin=(0, 0, 0, 0))
    c1 = card(s, rx, 2.75, rw, 1.45, fill=WHITE, line=RULE, name="CL02原文占位")
    box_text(c1, [[("CNAS-CL02:2023　7.3.7.4 检验结果的可比性", {"bold": True, "color": NAVY})],
                  [("【待核：从CL02:2023 / ISO 15189:2022正版文本逐字复制2–3句】", {"color": RED})]],
             size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE, margin=(0.2, 0.08, 0.15, 0.08))
    c2 = card(s, rx, 4.4, rw, 2.25, fill=CARD, name="A001原文")
    box_text(c2, [[("CNAS-CL02-A001:2023　7.3.7.4 1)", {"bold": True, "color": NAVY})],
                  [("实验室内部结果比对的程序文件应规定", {}), ("比对条件、样品类型及数量、比对方案、判断标准、频次及相关措施", {"bold": True, "color": TEAL}),
                   ("，可参考 CNAS-GL047 以及相关国家/行业标准，如 WS/T 406、WS/T 407。", {})]],
             size=16, color=INK, anchor=MSO_ANCHOR.TOP, margin=(0.2, 0.12, 0.18, 0.08))
    set_notes(s, """【CNAS-CL02:2023 前言（CNAS官网PDF，已核）】“本准则等同采用ISO 15189:2022《医学实验室—质量和能力的要求》……本准则内容及条款号与国际标准ISO 15189:2022《医学实验室—质量和能力的要求》内容及条款号完全一致。CNAS鼓励申请认可的机构购买和使用正版ISO标准及正版国家标准。ISO 15189:2022及本准则均受到版权保护，未经恰当的授权禁止复制。……本准则为第四版，取代CNAS-CL02:2012(第三版)。”2023年6月1日发布，2023年12月1日实施。CNAS《实验室认可规范文件清单》（2026-9-23）仍列为有效。
CL02:2023 第7.3.7.4条正文：CNAS官网不公开（版权），【待核：从正版文本抄录】。

【CNAS-CL02-A001:2023 第7.3.7.4条原文（第15页，已核）】
7.3.7.4 检验结果可比性　应符合 ISO 15189，7.3.7.4 条款及下列要求：
1) 实验室内部结果比对的程序文件应规定比对条件、样品类型及数量、比对方案、判断标准、频次及相关措施，可参考 CNAS-GL047 以及相关国家/行业标准，如 WS/T 406、WS/T 407。
2) 应规定由多个人员进行的手工检验项目比对的方法和判断标准，例如：显微镜检查、培养结果判读、抑菌圈测量等，定期（至少每 6 个月 1 次，每次至少 5 份临床样品）进行检验人员的结果比对。
3) 比对记录应由授权人员审核并签字，并至少保留 2 年。
说明：A001没有规定多系统定量比对的具体频次和样本数（由程序文件自行规定）；A001全文未出现“POCT”。

【原稿引用的CL02:2012 5.6.4（已废止，仅供对照）】应规定比较程序和所用设备和方法，以及建立临床适宜区间内患者样品结果可比性的方法。此要求适用于相同或不同的程序、设备、不同地点或所有这些情况。注：在测量结果可溯源至同一标准的特定情况下，如校准物可互换，则认为结果具有计量学可比性。……
左侧为原稿的2012版封面截图，待替换为2023版封面。""")


def p09_pyramid(prs, s, ctx):
    chrome(s, T1, "文件层级：从规定到我科程序", "各文件原文 · CNAS-CL02-A001:2023 §7.3.7.4", badge_w=4.4)
    layers = [("部门规范性文件", NAVY, WHITE), ("行业标准", RGBColor(0x2C, 0x5A, 0x86), WHITE), ("认可准则", TEAL, WHITE),
              ("认可指南", RGBColor(0x6D, 0xBA, 0xB5), NAVY), ("我科程序文件", CARD, NAVY)]
    docs = ["《医疗机构临床实验室管理办法》\n《医疗机构检查检验结果互认管理办法》",
            "WS/T 407-2012　·　WS/T 885—2026",
            "CNAS-CL02:2023 第7.3.7.4条\nCNAS-CL02-A001:2023 第7.3.7.4条",
            "CNAS-GL047:2021",
            "LZJY-2-CX-38《比对管理程序》"]
    top, H, base_w, top_w, cx = 1.7, 4.9, 5.6, 2.3, MX + 2.8
    lh = H / 5
    for i, ((name, fill, tc), doc) in enumerate(zip(layers, docs)):
        y = top + i * lh
        w_top = top_w + (base_w - top_w) * i / 5; w_bot = top_w + (base_w - top_w) * (i + 1) / 5
        gapy = 0.06
        shp = s.shapes.add_shape(MSO_SHAPE.TRAPEZOID, Inches(cx - w_bot / 2), Inches(y), Inches(w_bot), Inches(lh - gapy))
        inset = (w_bot - w_top) / 2
        shp.adjustments[0] = inset / min(w_bot, lh - gapy)
        shp.name = f"金字塔层{i+1}"
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
        if fill == CARD:
            shp.line.color.rgb = NAVY; shp.line.width = Pt(1.25)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_text(tf, [[(f"{'①②③④⑤'[i]} {name}", {"bold": True})]], size=15, color=tc, align=PP_ALIGN.CENTER)
        add_text(s, 6.95, y + 0.05, SLIDE_W - MX - 6.95, lh - 0.15, [[(l, {"bold": i == 4})] for l in doc.split("\n")], size=16,
                 color=NAVY if i == 4 else INK, anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0, 0, 0))
        arrow(s, cx + (w_top + w_bot) / 4 + 0.1, y + lh / 2 - 0.03, 6.85, y + lh / 2 - 0.03, color=RULE, width=1.25, head=False)
    # scope brackets
    add_text(s, 6.95, 6.62, 5.8, 0.35, [[("①②：所有医疗机构　｜　③④：申请或已获CNAS认可的实验室", {})]], size=13, color=MUTED, margin=(0, 0, 0, 0))
    set_notes(s, """【层级说明】
① 部门规范性文件：卫医发〔2006〕73号《医疗机构临床实验室管理办法》（第24、29、30条）；国卫医发〔2022〕6号《医疗机构检查检验结果互认管理办法》。
② 卫生行业标准（推荐性）：WS/T 407-2012（现行，适用于同一医疗机构内的临床实验室）；WS/T 885—2026（2026-11-01实施，4.2.3条要求实验室内部比对遵循WS/T 407等）。
③ 认可准则：CNAS-CL02:2023 第7.3.7.4条（等同ISO 15189:2022）；CNAS-CL02-A001:2023 第7.3.7.4条——明确“可参考 CNAS-GL047 以及相关国家/行业标准，如 WS/T 406、WS/T 407”。
④ 认可指南：CNAS-GL047:2021（适用于申请认可或已获认可的医学实验室；指南文件，可参考使用）。
⑤ 我科：LZJY-2-CX-38《比对管理程序》（第5/2版，2022-08-30发布，2022-09-01实施）。
适用对象：①②面向医疗机构；③④面向申请或已获CNAS认可的实验室（GL047“1 范围”原文）。""")
