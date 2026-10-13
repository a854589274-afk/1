# -*- coding: utf-8 -*-
"""在原 PPT 上就地改造：读取原文段落（保留红字/上下标），清空原页后按网格重排；插入新页并重排顺序。"""
import io, os, copy
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree
from lib import *

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, '..')
IMG = os.path.join(ROOT, '插图')
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
COLMAP = {'FF0000': C['RED'], 'C00000': C['RED'], '0070C0': C['NAVY']}


def extract(prs):
    """返回 O[页][shape_id] = [段落...]；PICS[页] = [(id, blob, crop, w, h)]"""
    O, PICS = {}, {}
    for i, s in enumerate(prs.slides, 1):
        O[i], PICS[i] = {}, []
        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6:
                    walk(sh.shapes); continue
                if sh.shape_type == 13:
                    PICS[i].append((sh.shape_id, sh.image.blob,
                                    (sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom),
                                    sh.width / EMU, sh.height / EMU))
                if sh.has_text_frame:
                    paras = []
                    for p in sh.text_frame.paragraphs:
                        segs = []
                        for r in p.runs:
                            rPr = r._r.find(A + 'rPr')
                            d = dict(t=r.text)
                            if rPr is not None:
                                if rPr.get('baseline') not in (None, '0'):
                                    d['base'] = int(rPr.get('baseline'))
                                sf = rPr.find(A + 'solidFill')
                                if sf is not None and sf[0].get('val') in COLMAP:
                                    d['color'] = COLMAP[sf[0].get('val')]
                            if d['t']:
                                segs.append(d)
                        if ''.join(x['t'] for x in segs).strip():
                            # 去首尾空白
                            while segs and not segs[0]['t'].strip(): segs.pop(0)
                            while segs and not segs[-1]['t'].strip(): segs.pop()
                            segs[0]['t'] = segs[0]['t'].lstrip(); segs[-1]['t'] = segs[-1]['t'].rstrip()
                            paras.append(segs)
                    if paras:
                        O[i][sh.shape_id] = paras
        walk(s.shapes)
    return O, PICS


def clear_slide(slide):
    tree = slide.shapes._spTree
    for el in list(tree):
        if el.tag in (qn('p:sp'), qn('p:pic'), qn('p:grpSp'), qn('p:graphicFrame'), qn('p:cxnSp')):
            tree.remove(el)


class Deck:
    def __init__(self, src, mode):
        self.prs = Presentation(src)
        self.mode = mode
        self.O, self.PICS = extract(self.prs)
        self.orig_slides = list(self.prs.slides)
        self.order = []        # 终稿顺序（slide 对象）
        self.log = {}          # 页 -> dict(template, added, replaced, pending, orig)
        self.layout = self.prs.slide_layouts[6] if len(self.prs.slide_layouts) > 6 else self.prs.slide_layouts[0]
        for sl in self.orig_slides:
            clear_slide(sl)
            for ph in list(sl.placeholders):
                ph._element.getparent().remove(ph._element)

    def orig(self, n):
        sl = self.orig_slides[n - 1]
        self.order.append(sl)
        return sl

    def new(self):
        sl = self.prs.slides.add_slide(self.layout)
        for ph in list(sl.placeholders):
            ph._element.getparent().remove(ph._element)
        self.order.append(sl)
        return sl

    def T(self, n, sid, k=None):
        """取原文段落（深拷贝）。"""
        ps = copy.deepcopy(self.O[n][sid])
        return ps if k is None else ps[k]

    def pic(self, n, sid):
        for p in self.PICS[n]:
            if p[0] == sid:
                return p
        raise KeyError((n, sid))

    def finish(self, out):
        lst = self.prs.slides._sldIdLst
        ids = {sl.slide_id: el for el in lst for sl in self.prs.slides if el.get('id') == str(sl.slide_id)}
        els = []
        for sl in self.order:
            for el in lst:
                if el.get('id') == str(sl.slide_id):
                    els.append(el)
        for el in list(lst): lst.remove(el)
        for el in els: lst.append(el)
        # 主题字体统一为微软雅黑/Arial
        for m in self.prs.slide_masters:
            for rel in m.part.rels.values():
                if 'theme' in rel.reltype:
                    th = etree.fromstring(rel.target_part.blob)
                    for tag in ('majorFont', 'minorFont'):
                        for f in th.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}' + tag):
                            f.find(A + 'latin').set('typeface', FONT_LAT)
                            f.find(A + 'ea').set('typeface', FONT_EA)
                            for fn in f.findall(A + 'font'):
                                f.remove(fn)
                    rel.target_part._blob = etree.tostring(th, xml_declaration=True, encoding='UTF-8', standalone=True)
        self.prs.save(out)


def img(name):
    return os.path.join(IMG, name + '.png')
