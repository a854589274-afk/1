# -*- coding: utf-8 -*-
"""原生图表（可编辑）：环形图、柱状图。"""
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from lib import EMU, C, FONT_LAT

def doughnut(sl, x, y, d, vals, hi=None, cats=None):
    cd = CategoryChartData()
    cd.categories = cats or [str(i + 1) for i in range(len(vals))]
    cd.add_series('条数', vals)
    gf = sl.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Emu(int(x * EMU)), Emu(int(y * EMU)), Emu(int(d * EMU)), Emu(int(d * EMU)), cd)
    gf.name = 'CHART'
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    from lxml import etree
    from pptx.oxml.ns import qn
    hs = plot._element.find(qn('c:holeSize'))
    if hs is not None: hs.set('val', '58')
    cols = ['C5D0DC', 'A9B8C9', C['ORANGE'], '8EA2B8', '7590AD', '5B6B7F']
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(C['ORANGE'] if i == hi else ['1F3A5F', '2E5482', '3F6A9A', '5B7FA8', '7C98BB', 'A3B6CE'][i])
        pt.format.line.color.rgb = RGBColor.from_string('FFFFFF')
    return gf

def bars(sl, x, y, w, h, cats, vals, colors, ymax=None, number_format='0.00"%"', label_size=12):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series('值', vals)
    gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(int(x * EMU)), Emu(int(y * EMU)), Emu(int(w * EMU)), Emu(int(h * EMU)), cd)
    gf.name = 'CHART'
    ch = gf.chart
    ch.has_legend = False; ch.has_title = False
    ch.font.size = Pt(label_size); ch.font.name = FONT_LAT
    va = ch.value_axis
    if ymax is not None:
        va.maximum_scale = ymax; va.minimum_scale = 0
    va.has_major_gridlines = False
    va.tick_labels.font.size = Pt(label_size)
    va.format.line.color.rgb = RGBColor.from_string('C9D1DA')
    ch.category_axis.tick_labels.font.size = Pt(label_size)
    ch.category_axis.format.line.color.rgb = RGBColor.from_string('C9D1DA')
    plot = ch.plots[0]
    plot.gap_width = 80
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(label_size); dl.font.bold = True
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid(); pt.format.fill.fore_color.rgb = RGBColor.from_string(colors[i])
    return gf
