# -*- coding: utf-8 -*-
"""质检：①文字守恒核对（原PPT逐段 → 标注版/清洁版）②版式检查（溢出、重叠、越界、新图数量、页数）。
用法：python3 qa.py 原PPT.pptx 输出目录
文字溢出用 LibreOffice 独立复核：把每个文本框改为"根据文字调整形状"，让 LibreOffice 重新排版后读回高度，与原高度比较。"""
import sys, os, re, json, copy, subprocess, tempfile, shutil, itertools, csv
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

EMU = 914400
SW, SH, ML, MR, MT, MB = 13.333, 7.5, 0.6, 0.6, 0.45, 0.45
WS = re.compile(r'[\s　  ​]+')
norm = lambda s: WS.sub('', s)


# ---------------- 文字抽取 ----------------
def para_texts(shape):
    out = []
    if shape.shape_type == 6:
        for s in shape.shapes: out += para_texts(s)
        return out
    if getattr(shape, 'has_table', False) and shape.has_table:
        for r in shape.table.rows:
            for c in r.cells:
                out += [p.text for p in c.text_frame.paragraphs]
        return out
    if shape.has_text_frame:
        out += [''.join(r.text for r in p.runs) for p in shape.text_frame.paragraphs]
    return out


def slide_text(sl):
    t = []
    for sh in sl.shapes: t += para_texts(sh)
    return norm(''.join(t))


def greedy(p, S):
    """p 是否能由 S 中的若干片段依序拼出（每段≥2字，最后一段可1字）。返回片段数或 None。"""
    if not p: return 0
    if p in S: return 1
    i, n, pos = 0, 0, 0
    while i < len(p):
        best = 0
        lo, hi = 1, len(p) - i
        while lo <= hi:                       # 二分找最长可匹配前缀
            mid = (lo + hi) // 2
            if p[i:i + mid] in S: best = mid; lo = mid + 1
            else: hi = mid - 1
        if best == 0 or (best < 2 and len(p) - i > 1): return None
        i += best; n += 1
    return n


def conservation(orig_pptx, annot, clean, log):
    O = Presentation(orig_pptx)
    A, Cl = Presentation(annot), Presentation(clean)
    notes = log['annot']['notes']
    replog = log['annot']['replog']
    m = {}
    for pg, d in notes.items():
        for n in re.findall(r'原P(\d+)', d.get('orig', '')):
            m.setdefault(int(n), []).append(int(pg))
    rows, per = [], {}
    for n, sl in enumerate(O.slides, 1):
        pages = sorted(set(m.get(n, [])))
        SA = ''.join(slide_text(A.slides[p - 1]) for p in pages)
        SC = ''.join(slide_text(Cl.slides[p - 1]) for p in pages)
        paras = []
        for sh in sl.shapes:
            paras += [t for t in para_texts(sh) if norm(t)]
        stat = dict(total=0, same=0, split=0, rep=0, missing=0)
        for t in paras:
            p = norm(t)
            stat['total'] += 1
            here = [r for r in replog if int(r[0]) in pages and ((r[2] and norm(r[2]) in p) or (not r[2] and norm(r[5]) in p))]
            if here:
                exp = p
                for r in here:
                    if r[2]: exp = exp.replace(norm(r[2]), norm(r[3]), 1)
                    else: exp = exp.replace(norm(r[5]), norm(r[3]) + norm(r[5]), 1)
                ok_clean = greedy(exp, SC)
                okA = all((not r[2] or norm(r[2]) in SA) and norm(r[3]) in SA for r in here) and greedy(p, SA.replace('［', '').replace('］', '')) is not None
                codes = '、'.join(sorted(set(r[1] for r in here)))
                status = f'替换段（{codes}）：' + ('清洁版新文完整' if ok_clean else '清洁版未找到') + '；' + ('标注版旧文与新文均在' if okA else '标注版缺旧文或新文')
                stat['rep'] += 1
                if not ok_clean or not okA: stat['missing'] += 1
            else:
                g = greedy(p, SA); gc = greedy(p, SC)
                if g == 1 and gc == 1:
                    status = '原样出现'; stat['same'] += 1
                elif g and gc:
                    status = f'拆分出现（{max(g, gc)}段，字符序列完整）'; stat['split'] += 1
                else:
                    status = '缺失'; stat['missing'] += 1
            rows.append(dict(原页=n, 终稿页=','.join(map(str, pages)), 原段落=t.strip()[:60] + ('…' if len(t.strip()) > 60 else ''), 结果=status))
        per[n] = stat
    return rows, per


# ---------------- 版式检查 ----------------
def boxes(sl):
    out = []
    for sh in sl.shapes:
        if sh.left is None: continue
        out.append((sh.name, sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU, sh))
    return out


def overlap(a, b):
    ax0, ay0, ax1, ay1 = a[1], a[2], a[1] + a[3], a[2] + a[4]
    bx0, by0, bx1, by1 = b[1], b[2], b[1] + b[3], b[2] + b[4]
    w = min(ax1, bx1) - max(ax0, bx0); h = min(ay1, by1) - max(ay0, by0)
    return w * h if w > 0.01 and h > 0.01 else 0


SKIP_OV = ('LINE', 'ARROW', 'DOT', 'NEWIMG_HERO', 'SEG')


def layout(deck, notes):
    prs = Presentation(deck)
    res = []
    for i, sl in enumerate(prs.slides, 1):
        bs = boxes(sl)
        hero = any(b[0] == 'NEWIMG_HERO' for b in bs)
        oob = [b[0] for b in bs if b[0] != 'NEWIMG_HERO' and (b[1] < ML - 0.01 or b[2] < MT - 0.01 or b[1] + b[3] > SW - MR + 0.01 or b[2] + b[4] > SH - MB + 0.01)]
        ovs = []
        cand = [b for b in bs if not b[0].startswith(SKIP_OV)]
        for a, b in itertools.combinations(cand, 2):
            if overlap(a, b) > 0: ovs.append(f'{a[0]}×{b[0]}')
        imgs = notes.get(str(i), {}).get('images', [])
        groups = sorted(set(re.sub(r'-\d+$', '', re.sub(r'c$', '', x)) for x in imgs))
        res.append(dict(page=i, out_of_bounds=oob, overlaps=ovs, new_image_groups=groups))
    return res, len(prs.slides)


def lo_overflow(deck):
    """LibreOffice 独立复核文字溢出：所有自动换行文本框改为 spAutoFit，经 LO 重排后比较高度。"""
    tmp = tempfile.mkdtemp()
    prs = Presentation(deck)
    rec = []
    for i, sl in enumerate(prs.slides, 1):
        k = 0
        for sp in sl.shapes._spTree.iter(qn('p:sp')):
            tx = sp.find(qn('p:txBody'))
            if tx is None: continue
            bp = tx.find(qn('a:bodyPr'))
            txt = ''.join(t.text or '' for t in tx.iter(qn('a:t')))
            if not txt.strip() or bp.get('wrap') == 'none': continue
            nm = sp.find(qn('p:nvSpPr')).find(qn('p:cNvPr'))
            nm.set('name', f'QA{i}_{k}')
            xfrm = sp.find(qn('p:spPr')).find(qn('a:xfrm'))
            h = int(xfrm.find(qn('a:ext')).get('cy')) / EMU
            for e in list(bp): bp.remove(e)
            etree.SubElement(bp, qn('a:spAutoFit'))
            rec.append((i, f'QA{i}_{k}', h, txt[:30]))
            k += 1
    src = os.path.join(tmp, 'qa.pptx'); prs.save(src)
    od = os.path.join(tmp, 'out'); os.makedirs(od)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pptx:Impress MS PowerPoint 2007 XML', '--outdir', od, src],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    P2 = Presentation(os.path.join(od, 'qa.pptx'))
    hs = {}
    for i, sl in enumerate(P2.slides, 1):
        for sp in sl.shapes._spTree.iter(qn('p:sp')):
            nm = sp.find(qn('p:nvSpPr')).find(qn('p:cNvPr')).get('name')
            xfrm = sp.find(qn('p:spPr')).find(qn('a:xfrm'))
            if xfrm is not None and nm.startswith('QA'):
                hs[nm] = int(xfrm.find(qn('a:ext')).get('cy')) / EMU
    shutil.rmtree(tmp)
    over = [(i, nm, round(h, 2), round(hs[nm], 2), t) for i, nm, h, t in rec if nm in hs and hs[nm] > h + 0.03]
    return over, len(rec), len(hs)


def fontcheck(deck):
    prs = Presentation(deck)
    small, fonts = [], set()
    for i, sl in enumerate(prs.slides, 1):
        for r in sl.shapes._spTree.iter(qn('a:rPr')):
            sz = r.get('sz')
            t = r.getparent().find(qn('a:t'))
            txt = (t.text if t is not None else '') or ''
            if sz and int(sz) < 1200 and txt.strip():
                small.append((i, int(sz) / 100, txt[:20]))
            for tag in ('a:latin', 'a:ea'):
                e = r.find(qn(tag))
                if e is not None: fonts.add(e.get('typeface'))
    return small, fonts


if __name__ == '__main__':
    orig, outdir = sys.argv[1], sys.argv[2]
    A = os.path.join(outdir, '定量检验程序结果可比性验证_标注版.pptx')
    Cl = os.path.join(outdir, '定量检验程序结果可比性验证_清洁版.pptx')
    log = json.load(open(os.path.join(outdir, '_build_log.json')))
    rows, per = conservation(orig, A, Cl, log)
    with open(os.path.join(outdir, '文字守恒核对表.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
    report = dict(per_slide=per)
    for name, deck in (('annot', A), ('clean', Cl)):
        lay, n = layout(deck, log[name]['notes'])
        over, nrec, nret = lo_overflow(deck)
        small, fonts = fontcheck(deck)
        report[name] = dict(pages=n, layout=lay, lo_overflow=over, lo_checked=nrec, lo_returned=nret, small_fonts=small, fonts=sorted(fonts))
    json.dump(report, open(os.path.join(outdir, '_qa.json'), 'w'), ensure_ascii=False, indent=1)
    tot = {k: sum(v[k] for v in per.values()) for k in ('total', 'same', 'split', 'rep', 'missing')}
    print('守恒', tot)
    for name in ('annot', 'clean'):
        r = report[name]
        print(name, '页数', r['pages'], 'LO溢出', len(r['lo_overflow']), '/', r['lo_checked'], '越界页', [x['page'] for x in r['layout'] if x['out_of_bounds']],
              '重叠页', [x['page'] for x in r['layout'] if x['overlaps']], '无新图页', [x['page'] for x in r['layout'] if not x['new_image_groups']],
              '新图>2组', [x['page'] for x in r['layout'] if len(x['new_image_groups']) > 2], '小于12pt', len(r['small_fonts']), '字体', r['fonts'])
